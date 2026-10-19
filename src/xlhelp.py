"""Utilidades openpyxl con la presentación de los anexos del modelo (encabezado de proyecto, celeste, bordes finos)."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.worksheet.page import PageMargins
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).parent))
from params import PROYECTO

NOMBRE = PROYECTO["nombre"]
UBIC = {"departamento": PROYECTO["departamento"].upper(), "provincia": PROYECTO["provincias"].upper(), "distrito": PROYECTO["distritos"].upper()}
PROY = PROYECTO["proyectista"]
FECHA = PROYECTO["fecha"].upper()
AZ = PatternFill("solid", fgColor="C5DFF0")    # celeste del modelo
AZ2 = PatternFill("solid", fgColor="DDEBF7")
AM = PatternFill("solid", fgColor="FFF2CC")    # celdas de ingreso de datos
VE = PatternFill("solid", fgColor="E2EFDA")
thin = Side(style="thin", color="000000")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
F = Font(name="Arial", size=9)
FB = Font(name="Arial", size=9, bold=True)
FT = Font(name="Arial", size=11, bold=True)
FI = Font(name="Arial", size=8, italic=True, color="555555")
FBL = Font(name="Arial", size=9, color="0000FF")      # dato de entrada
CEN = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="center")


def libro():
    wb = Workbook(); wb.remove(wb.active); return wb


def hoja(wb, titulo, anchos, apaisado=False):
    ws = wb.create_sheet(titulo)
    for i, w in enumerate(anchos, 1):
        ws.column_dimensions[CL(i)].width = w
    ws.sheet_view.showGridLines = False
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape" if apaisado else "portrait"
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.5, right=0.4, top=0.6, bottom=0.6)
    ws.oddFooter.left.text = "Proyectista: " + PROY; ws.oddFooter.left.size = 7
    ws.oddFooter.right.text = "Pág. &P de &N"; ws.oddFooter.right.size = 7
    return ws


def cab(ws, titulo, subtitulo=None, ncol=8, fila=1):
    """Encabezado tipo modelo: nombre del proyecto en banda celeste + título del cálculo."""
    last = CL(ncol)
    ws.merge_cells(f"A{fila}:{last}{fila}"); c = ws[f"A{fila}"]
    c.value = NOMBRE; c.font = Font(name="Arial", size=8, bold=True); c.alignment = CEN; c.fill = AZ
    ws.row_dimensions[fila].height = 30
    ws.merge_cells(f"A{fila+1}:{last}{fila+1}"); c = ws[f"A{fila+1}"]
    c.value = titulo; c.font = FT; c.alignment = CEN
    ws.row_dimensions[fila + 1].height = 20
    r = fila + 2
    if subtitulo:
        ws.merge_cells(f"A{r}:{last}{r}"); c = ws[f"A{r}"]; c.value = subtitulo; c.font = FB; c.alignment = CEN; r += 1
    return r + 1


def put(ws, ref, v, font=F, fill=None, fmt=None, al=None, border=True):
    c = ws[ref]; c.value = v; c.font = font
    if fill: c.fill = fill
    if fmt: c.number_format = fmt
    c.alignment = al or (LEFT if isinstance(v, str) and not str(v).startswith("=") else CEN)
    if border: c.border = BOX
    return c


def fila(ws, r, valores, col0=1, font=F, fill=None, fmts=None, border=True):
    for k, v in enumerate(valores):
        if v is None and not border: continue
        c = ws.cell(r, col0 + k); c.value = v; c.font = font
        if fill: c.fill = fill
        if fmts and k < len(fmts) and fmts[k]: c.number_format = fmts[k]
        c.alignment = LEFT if (isinstance(v, str) and not v.startswith("=") and k == 0) else CEN
        if border: c.border = BOX


def titulo_seccion(ws, r, texto, ncol=8, fill=AZ2):
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncol)
    c = ws.cell(r, 1); c.value = texto; c.font = FB; c.alignment = LEFT
    if fill: c.fill = fill
    for k in range(1, ncol + 1): ws.cell(r, k).border = BOX


def nota(ws, r, texto, ncol=8, h=None):
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncol)
    c = ws.cell(r, 1); c.value = texto; c.font = FI; c.alignment = LEFT
    if h: ws.row_dimensions[r].height = h


def dato(ws, r, desc, simbolo, valor, unidad, fuente="", col0=1, nombre=None, fmt=None, entrada=True):
    """Fila de dato: descripción | símbolo | valor | unidad | fuente. Devuelve la referencia absoluta del valor."""
    vals = [desc, simbolo, valor, unidad, fuente]
    for k, v in enumerate(vals):
        c = ws.cell(r, col0 + k); c.value = v; c.border = BOX
        c.font = (FBL if (k == 2 and entrada and not (isinstance(v, str) and v.startswith("="))) else F)
        c.alignment = LEFT if k in (0, 4) else CEN
        if k == 2:
            if fmt: c.number_format = fmt
            if entrada and not (isinstance(v, str) and v.startswith("=")): c.fill = AM
    ref = f"${CL(col0+2)}${r}"
    return ref
