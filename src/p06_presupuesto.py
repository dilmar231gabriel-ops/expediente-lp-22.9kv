"""Volumen V – Metrado y presupuesto con la organización de hojas del modelo Ayaviri–Pucará:
RESUMEN OBRA, GG, GS, CCL, CRONOG_VALORIZ, CRONOG_SUMINISTRO, GANTT, SUMINISTRO LP, MONTAJE LP, TRANSPORTE LP, APU,
FP, PLANILLA (inventario físico por estructura), BOM_ARMADOS, BP, MO, EQUIPO, MATERIALES.
Todo valor derivado es fórmula. Uso: python p06_presupuesto.py WORK SALIDA.xlsx"""
import json, math, sys
from pathlib import Path
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as CL
sys.path.insert(0, str(Path(__file__).parent))
from params import *
from catalogo import *
from xlhelp import libro, hoja, put, fila, titulo_seccion, nota, BOX, F, FB, FT, FI, FBL, AZ, AZ2, AM, VE, CEN, LEFT, NOMBRE, PROY

W = Path(sys.argv[1]); OUT = Path(sys.argv[2])
E = pd.read_csv(W / "planilla_estructuras.csv"); R = json.load(open(W / "resultados.json"))
N = len(E); Lkm = R["plantillado"]["L_troncal_m"] / 1000
PLAZO = PROYECTO["plazo_dias"]
MON = '#,##0.00;-#,##0.00;"-"'
UBIC = PROYECTO["ubicacion"]
wb = libro()

def encab(ws, titulo, sub=None, ncol=10):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    c = ws.cell(1, 1, "VALOR REFERENCIAL"); c.font = FB
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncol)
    c = ws.cell(2, 1, titulo); c.font = FT
    ws.cell(3, 1, "PROYECTO").font = FB; ws.cell(3, 3, ":"); ws.cell(3, 4, NOMBRE).font = F
    ws.cell(4, 1, "UBICACIÓN").font = FB; ws.cell(4, 3, ":"); ws.cell(4, 4, UBIC).font = F
    if sub: ws.cell(5, 1, sub).font = FB
    return 7

# =============================================================================== BP (base de precios)
bp = hoja(wb, "BP", [10, 70, 8, 12, 12, 14, 34])
r = encab(bp, "BASE DE PRECIOS: SUMINISTROS, MANO DE OBRA, EQUIPOS E INSUMOS DE MONTAJE", ncol=7)
put(bp, f"B{r}", "Tipo de cambio venta (S/ por US$)", font=FB); put(bp, f"F{r}", TC, font=FBL, fill=AM, fmt="0.000"); put(bp, f"G{r}", TC_FUENTE, font=FI)
TCREF = f"BP!$F${r}"; r += 2
fila(bp, r, ["Código", "Descripción", "Und.", "Precio US$", "Precio S/ (dato)", "Precio S/ (sin IGV)", "Fuente"], font=FB, fill=AZ); r += 1
BPROW = {}
for g, gname in GRUPOS:
    titulo_seccion(bp, r, f"{g}  {gname}", ncol=7); r += 1
    for m in MAT:
        if m[0].split(".")[0] != g: continue
        cod, desc, und, usd, sol, fte, _ = m
        if cod in PRECIOS_PROYECTO: usd, sol, fte = None, PRECIOS_PROYECTO[cod], "PROY"
        fila(bp, r, [cod, desc, und, usd, sol, f"=IF(D{r}<>\"\",ROUND(D{r}*{TCREF},2),E{r})",
                     {"COT-M": "Cotización exp. de referencia 2022 (US$) × TC vigente – ACTUALIZAR", "COT-MS": "Precio exp. de referencia 2022 (S/) – ACTUALIZAR",
                      "EST": "Estimado del proyectista – COTIZAR", "PROY": "Precio del proyecto (proyecto.yaml → precios)"}[fte]], fmts=[None, None, None, "0.0000", MON, MON, None])
        for cc in "DE":
            if bp[f"{cc}{r}"].value is not None: bp[f"{cc}{r}"].fill = AM; bp[f"{cc}{r}"].font = FBL
        BPROW[cod] = r; r += 1
r += 1
titulo_seccion(bp, r, "MANO DE OBRA (S/ por hora-hombre, incluye leyes sociales)", ncol=7); r += 1
MOROW = {}
for k, (desc, (pr, fte)) in enumerate(MO.items()):
    fila(bp, r, [f"MO{k+1:02d}", desc, "h-h", None, pr, f"=E{r}", fte], fmts=[None, None, None, None, MON, MON, None])
    bp[f"E{r}"].fill = AM; bp[f"E{r}"].font = FBL; MOROW[desc] = r; r += 1
r += 1
titulo_seccion(bp, r, "EQUIPOS (S/ por hora-máquina, incluye operador y combustible cuando corresponde)", ncol=7); r += 1
EQROW = {}
for k, (desc, (pr, fte)) in enumerate(EQ.items()):
    fila(bp, r, [f"EQ{k+1:02d}", desc, "h-m", None, pr, f"=E{r}", fte], fmts=[None, None, None, None, MON, MON, None])
    bp[f"E{r}"].fill = AM; bp[f"E{r}"].font = FBL; EQROW[desc] = r; r += 1
r += 1
titulo_seccion(bp, r, "INSUMOS DE MONTAJE", ncol=7); r += 1
MMROW = {}
for k, (desc, (und, pr, fte)) in enumerate(MATMON.items()):
    fila(bp, r, [f"IM{k+1:02d}", desc, und, None, pr, f"=E{r}", fte], fmts=[None, None, None, None, MON, MON, None])
    bp[f"E{r}"].fill = AM; bp[f"E{r}"].font = FBL; MMROW[desc] = r; r += 1
r += 1
nota(bp, r, f"Mano de obra de construcción: jornal básico del convenio CAPECO–FTCCP 2025-2026 (Operario S/ {JORNAL_CAPECO['Operario']:.2f}, Oficial "
     f"S/ {JORNAL_CAPECO['Oficial']:.2f}, Peón S/ {JORNAL_CAPECO['Peón']:.2f}) × factor {FACTOR_HH} (leyes sociales y beneficios, SUPUESTO) / 8 h. "
     "Profesionales: tarifas de la base de precios del expediente de referencia LP Ayaviri–Pucará (2022). Equipos: tarifas de ACU de Electro Sur Este S.A.A. (2021). "
     "Todos los precios son sin IGV y deben actualizarse a la fecha de convocatoria.", ncol=7, h=55)

# =============================================================================== BOM_ARMADOS
bo = hoja(wb, "BOM_ARMADOS", [9, 62, 6] + [9] * 10, apaisado=True)
r = encab(bo, "LISTA DE MATERIALES POR ARMADO, RETENIDA Y PUESTA A TIERRA (cantidad por unidad)", ncol=13)
cols_b = ARMADOS + ["RI", "PAT-1", "PAT-2"]
fila(bo, r, ["Código", "Descripción", "Und."] + cols_b, font=FB, fill=AZ); BOH = r; r += 1
usados = sorted({c for a in BOM.values() for c in a} | set(RET_RI) | {c for k in PAT_KIT.values() for c in k},
                key=lambda c: [int(x) for x in c.split(".")])
BOROW = {}
for cod in usados:
    m = MATD[cod]
    vals = [cod, m[1], m[2]] + [BOM[a].get(cod, None) for a in ARMADOS] + [RET_RI.get(cod, None)] + [PAT_KIT[p].get(cod, None) for p in ("PAT-1", "PAT-2")]
    fila(bo, r, vals); BOROW[cod] = r; r += 1
fila(bo, r, ["PTE", "Conductor AAAC 70 mm² para puentes y bajadas a equipos", "m"] + [PUENTE_M.get(a, None) for a in ARMADOS] + [None] * 3)
BOROW["PTE"] = r; r += 1
BOR0, BOR1 = BOH + 1, r - 1
BOM_RANGE = f"BOM_ARMADOS!$D${BOR0}:${CL(3+len(cols_b))}${BOR1}"; BOM_CODES = f"BOM_ARMADOS!$A${BOR0}:$A${BOR1}"
BOM_HDR = f"BOM_ARMADOS!$D${BOH}:${CL(3+len(cols_b))}${BOH}"
nota(bo, r + 1, "El cable de retenida no figura por unidad: su longitud depende de la altura del poste (ver PLANILLA, fila 'Cable por retenida'). "
     "Láminas de detalle LD-LP-01 a LD-LP-11 (Vol. VII) muestran cada armado con la misma lista.", ncol=13, h=30)

# =============================================================================== PLANILLA (inventario físico por estructura)
pl = hoja(wb, "PLANILLA", [9, 58, 6, 12] + [8.5] * N, apaisado=True)
r = encab(pl, "INVENTARIO FÍSICO POR ESTRUCTURA (METRADO DE SUMINISTROS)", ncol=12)
attrs = [("N.°", lambda x: int(x.N)), ("Estructura", lambda x: x.estructura), ("Armado", lambda x: x.armado), ("Poste", lambda x: x.poste),
         ("N.° retenidas", lambda x: int(x.n_retenidas)), ("Puesta a tierra", lambda x: x.puesta_tierra),
         ("Longitud del poste (m)", lambda x: POSTES[x.poste]["L"]), ("Vano adelante (m)", lambda x: 0 if pd.isna(x.vano_adelante_m) else round(x.vano_adelante_m, 2)),
         ("Progresiva (m)", lambda x: round(x.progresiva_m, 2))]
AR = {}
for k, (lab, fn) in enumerate(attrs):
    pl.cell(r, 2, lab).font = FB; pl.cell(r, 2).fill = AZ2; pl.cell(r, 2).border = BOX
    for j, x in enumerate(E.itertuples()):
        c = pl.cell(r, 5 + j, fn(x)); c.font = F; c.border = BOX; c.alignment = CEN
    AR[lab] = r; r += 1
CR = {}
for lab, f in [("Cable por retenida (m)", lambda c: f"=ROUND(({c}{AR['Longitud del poste (m)']}-({c}{AR['Longitud del poste (m)']}/10+0.6)-0.6)/COS(RADIANS({RET['angulo_deg']}))+2,1)"),
               ("Puente (m)", lambda c: f"=IFERROR(INDEX({BOM_RANGE},MATCH(\"PTE\",{BOM_CODES},0),MATCH({c}{AR['Armado']},{BOM_HDR},0)),0)")]:
    pl.cell(r, 2, lab).font = FB; pl.cell(r, 2).fill = AZ2; pl.cell(r, 2).border = BOX
    for j in range(N):
        c = CL(5 + j); cell = pl.cell(r, 5 + j, f(c)); cell.font = FI; cell.border = BOX; cell.number_format = "0.0"
    CR[lab] = r; r += 1
r += 1
fila(pl, r, ["Código", "Descripción", "Und.", "TOTAL"] + [None] * N, font=FB, fill=AZ); r += 1
PLROW = {}
LASTC = CL(4 + N)
for g, gname in GRUPOS:
    titulo_seccion(pl, r, f"{g}  {gname}", ncol=4); r += 1
    for m in MAT:
        cod = m[0]
        if cod.split(".")[0] != g: continue
        fila(pl, r, [cod, m[1], m[2], f"=SUM(E{r}:{LASTC}{r})"], fmts=[None, None, None, "#,##0.00"])
        pl[f"D{r}"].font = FB
        for j in range(N):
            c = CL(5 + j); parts = []
            if cod in ("1.01", "1.02", "1.03"):
                pst = {"1.01": "12/300", "1.02": "12/400", "1.03": "13/400"}[cod]
                parts.append(f'IF({c}${AR["Poste"]}="{pst}",1,0)')
            if cod == "3.01":
                parts.append(f"3*{c}${AR['Vano adelante (m)']}*(1+{DESPERDICIO_COND})+3*{c}${CR['Puente (m)']}")
            if cod in COMUN: parts.append(str(COMUN[cod]))
            if cod in BOROW and any(BOM[a].get(cod) for a in ARMADOS):
                parts.append(f"IFERROR(INDEX({BOM_RANGE},MATCH($A{r},{BOM_CODES},0),MATCH({c}${AR['Armado']},{BOM_HDR},0)),0)")
            if cod in RET_RI:
                parts.append(f"{c}${AR['N.° retenidas']}*{RET_RI[cod]}")
            if cod == "7.01":
                parts.append(f"{c}${AR['N.° retenidas']}*{c}${CR['Cable por retenida (m)']}")
            if any(cod in PAT_KIT[p] for p in PAT_KIT):
                parts.append(f'IF({c}${AR["Puesta a tierra"]}="PAT-1",{PAT_KIT["PAT-1"].get(cod,0)},IF({c}${AR["Puesta a tierra"]}="PAT-2",{PAT_KIT["PAT-2"].get(cod,0)},0))')
            cell = pl.cell(r, 5 + j, ("=" + "+".join(parts)) if parts else 0)
            cell.font = F; cell.number_format = '0.##;-0.##;"-"'; cell.border = BOX; cell.alignment = CEN
        PLROW[cod] = r; r += 1
pl.freeze_panes = "E8"
pl.page_setup.fitToWidth = 10; pl.print_title_cols = "A:D"; pl.page_setup.paperSize = pl.PAPERSIZE_A3

# =============================================================================== SUMINISTRO LP
su = hoja(wb, "SUMINISTRO LP", [8, 66, 7, 13, 13, 15, 3, 11, 9])
r = encab(su, "SUMINISTRO DE MATERIALES LÍNEA PRIMARIA", "A: SUMINISTRO DE MATERIALES", ncol=6)
fila(su, r, ["ITEM", "DESCRIPCIÓN DE PARTIDAS", "UNID.", "Metrado", "Cost. Unit. S/", "TOTAL S/"], font=FB, fill=AZ); r += 1
SUSUB = {}; SUROW = {}
for g, gname in GRUPOS:
    put(su, f"A{r}", g, font=FB); put(su, f"B{r}", gname, font=FB); r += 1; g0 = r
    for m in MAT:
        cod = m[0]
        if cod.split(".")[0] != g: continue
        fila(su, r, [cod, f"=PLANILLA!B{PLROW[cod]}", m[2], f"=PLANILLA!D{PLROW[cod]}", f"=BP!F{BPROW[cod]}", f"=ROUND(D{r}*E{r},2)"],
             fmts=[None, None, None, "#,##0.00", MON, MON])
        su[f"B{r}"].alignment = LEFT; SUROW[cod] = r; r += 1
    fila(su, r, [None, f"SUB-TOTAL {g}: {gname}", None, None, None, f"=SUM(F{g0}:F{r-1})"], font=FB, fmts=[None] * 5 + [MON])
    SUSUB[g] = r; r += 2
fila(su, r, [None, "TOTAL SUMINISTRO DE MATERIALES", None, None, None, "=" + "+".join(f"F{v}" for v in SUSUB.values())], font=FB, fill=VE, fmts=[None] * 5 + [MON])
SU_TOT = f"'SUMINISTRO LP'!F{r}"
nota(su, r + 2, "Metrados enlazados a la hoja PLANILLA (inventario físico por estructura); precios enlazados a la hoja BP.", ncol=6)

# =============================================================================== APU
ap = hoja(wb, "APU", [3, 58, 9, 8, 11, 12, 13, 8, 11, 12])
ap.print_area = "A1:H2000"
r = 2; ap.cell(r, 2, "ANÁLISIS DE PRECIOS UNITARIOS").font = FT; r += 1
ap.cell(r, 2, NOMBRE).font = FI; r += 1
ap.cell(r, 2, "Columnas I–J: metrado de la partida y cantidad total del insumo (para la relación de insumos).").font = FI; r += 2
# partidas: (item, descripción, und, rendimiento/día, MO {desc: cuadrilla}, EQ {desc: cuadrilla}, MAT [(desc, und, cant)], %MO materiales varios, metrado (fórmula))
c5 = lambda pst: f'COUNTIF(PLANILLA!$E${AR["Poste"]}:${LASTC}${AR["Poste"]},"{pst}")'
cA = lambda a: f'COUNTIF(PLANILLA!$E${AR["Armado"]}:${LASTC}${AR["Armado"]},"{a}")'
cP = lambda p: f'COUNTIF(PLANILLA!$E${AR["Puesta a tierra"]}:${LASTC}${AR["Puesta a tierra"]},"{p}")'
NRET = f'SUM(PLANILLA!$E${AR["N.° retenidas"]}:${LASTC}${AR["N.° retenidas"]})'
COND_KM = f"PLANILLA!D{PLROW['3.01']}/1000"
VEXC = {"12/300": 0.955, "12/400": 0.955, "13/400": 1.272}; VCIC = {"12/300": 0.754, "12/400": 0.741, "13/400": 1.021}
VSOL = {"12": 0.050, "13": 0.064}; VELI = {"12/300": 0.188, "12/400": 0.205, "13/400": 0.235}
VRET_E, VRET_R = 1.28, 1.23; VPAT = 1.31          # m³ (Anexo 7 y lámina PAT)
ING, TOP, DIB, OPE, OFI, PEO, CAP_ = ("Ingeniero especialista en Líneas Primarias", "Topógrafo operador de estación total",
                                      "Técnico especialista en dibujo por computadora", "Operario", "Oficial", "Peón", "Capataz")
CEM, HOR, PIE, AGU = ("Cemento Portland tipo I (bolsa 42,5 kg)", "Hormigón", "Piedra mediana", "Agua")
PART = [
    ("1", "OBRAS PRELIMINARES", None),
    ("1.01", "Replanteo topográfico, ubicación de estructuras e ingeniería de detalle de la línea primaria", "km", 2.0,
     {ING: 1, TOP: 1, DIB: 0.5, OFI: 1, PEO: 2}, {"Camioneta rural 4x4": 0.5, "Estación total y accesorios (incluye GPS)": 1, "Equipo de comunicación (par)": 1},
     [("Estacas, yeso, pintura y materiales de replanteo", "glb", 35.0)], 0, f"=ROUND({Lkm},3)"),
    ("1.02", "Estudio de coordinación de la protección y ajuste del reconectador", "Gbl", 1 / 6,
     {"Especialista en estudios de coordinación y protección": 1}, {"Equipo de cómputo y plotter": 1}, [], 0.10, "=1"),
    ("1.03", "Monitoreo arqueológico durante las excavaciones (Plan de Monitoreo Arqueológico)", "Gbl", 1 / 25,
     {"Arqueólogo": 1, PEO: 1}, {"Camioneta rural 4x4": 0.5}, [], 0.05, "=1"),
    ("1.04", "Programa de monitoreo ambiental (calidad de aire, ruido y suelo)", "Gbl", 1 / 15,
     {"Especialista ambiental": 1, OFI: 1}, {"Camioneta rural 4x4": 0.5}, [], 0.30, "=1"),
    ("1.05", "Programa de manejo de residuos sólidos", "Gbl", 1 / 15, {OFI: 1, PEO: 2}, {"Camioneta rural 4x4": 0.25}, [], 0.25, "=1"),
    ("1.06", "Talleres de información a la población y medidas preventivas", "Gbl", 1 / 4, {"Especialista ambiental": 1}, {"Camioneta rural 4x4": 0.5}, [], 0.30, "=1"),
    ("1.07", "Cartel de obra 3,60 × 2,40 m", "u", 1.0, {OPE: 1, PEO: 2}, {"Camioneta rural 4x4": 0.25},
     [("Gigantografía y bastidor de madera para cartel de obra 3,60 x 2,40 m", "u", 1.0)], 0, "=1"),
    ("2", "TRANSPORTE DE MATERIALES DE ALMACÉN DE OBRA A PUNTO DE INSTALACIÓN", None),
    ("2.01", "Transporte de poste de 12 m de almacén a punto de izaje", "u", 10.0, {OPE: 1, PEO: 4}, {"Camión grúa de 6 t": 1}, [], 0,
     f"={c5('12/300')}+{c5('12/400')}"),
    ("2.02", "Transporte de poste de 13 m de almacén a punto de izaje", "u", 9.0, {OPE: 1, PEO: 4}, {"Camión grúa de 6 t": 1}, [], 0, f"={c5('13/400')}"),
    ("2.03", "Transporte de armados, retenidas y puesta a tierra de almacén a punto de instalación", "Jgo.", 25.0, {OFI: 1, PEO: 2}, {"Camión de 4 t": 1}, [], 0, f"={N}"),
    ("2.04", "Transporte de conductor AAAC 70 mm² de almacén a punto de instalación", "km", 12.0, {OFI: 1, PEO: 2}, {"Camión de 4 t": 1}, [], 0, f"=ROUND({COND_KM},3)"),
    ("3", "INSTALACIÓN DE POSTES", None),
    ("3.01", "Excavación en terreno tipo I (limo-arcilloso) para poste, incluye hoyo de cimentación", "m3", 5.0, {OFI: 0.2, PEO: 2}, {}, [], 0,
     f"=ROUND({c5('12/300')}*{VEXC['12/300']}+{c5('12/400')}*{VEXC['12/400']}+{c5('13/400')}*{VEXC['13/400']},2)"),
    ("3.02", "Izaje, alineamiento, identificación y señalización de poste de 12 m", "u", 8.0, {OPE: 1, OFI: 1, PEO: 4}, {"Camión grúa de 6 t": 1, "Tirfor de 3 t": 1}, [], 0,
     f"={c5('12/300')}+{c5('12/400')}"),
    ("3.03", "Izaje, alineamiento, identificación y señalización de poste de 13 m", "u", 7.0, {OPE: 1, OFI: 1, PEO: 4}, {"Camión grúa de 6 t": 1, "Tirfor de 3 t": 1}, [], 0, f"={c5('13/400')}"),
    ("3.04", "Cimentación con concreto ciclópeo 1:10 + 30 % PM de poste de 12 m (0,75 m³)", "u", 8.0, {OPE: 1, OFI: 1, PEO: 3}, {"Mezcladora de concreto 9 p3": 0.5},
     [(CEM, "bls", 0.75 * 2.9), (HOR, "m3", 0.75 * 0.83), (PIE, "m3", 0.75 * 0.48), (AGU, "m3", 0.75 * 0.16)], 0, f"={c5('12/300')}+{c5('12/400')}"),
    ("3.05", "Cimentación con concreto ciclópeo 1:10 + 30 % PM de poste de 13 m (1,02 m³)", "u", 6.0, {OPE: 1, OFI: 1, PEO: 3}, {"Mezcladora de concreto 9 p3": 0.5},
     [(CEM, "bls", 1.02 * 2.9), (HOR, "m3", 1.02 * 0.83), (PIE, "m3", 1.02 * 0.48), (AGU, "m3", 1.02 * 0.16)], 0, f"={c5('13/400')}"),
    ("3.06", "Solado de concreto 1:12, e = 0,10 m, para poste", "u", 20.0, {OPE: 1, PEO: 2}, {}, [(CEM, "bls", 0.064 * 3.0), (HOR, "m3", 0.064 * 1.1), (AGU, "m3", 0.064 * 0.18)], 0, f"={N}"),
    ("3.07", "Eliminación de material excedente (DMT ≤ 1 km)", "m3", 12.0, {PEO: 2}, {"Camión de 4 t": 0.5}, [], 0,
     f"=ROUND({c5('12/300')}*{VELI['12/300']}+{c5('12/400')}*{VELI['12/400']}+{c5('13/400')}*{VELI['13/400']},2)"),
    ("4", "INSTALACIÓN DE RETENIDAS", None),
    ("4.01", "Excavación en terreno tipo I para retenida (0,80 × 0,80 × 2,00 m)", "m3", 5.0, {OFI: 0.2, PEO: 2}, {}, [], 0, f"=ROUND({NRET}*{VRET_E},2)"),
    ("4.02", "Instalación de retenida inclinada (RI), incluye bloque de anclaje y aislador", "u", 6.0, {OPE: 1, OFI: 1, PEO: 2}, {"Tirfor de 3 t": 1, "Camioneta rural 4x4": 0.25}, [], 0, f"={NRET}"),
    ("4.03", "Relleno y compactación con material propio seleccionado en capas de 0,20 m", "m3", 6.0, {OFI: 0.2, PEO: 2}, {"Compactador vibratorio tipo plancha": 0.5}, [(AGU, "m3", 0.05)], 0,
     f"=ROUND({NRET}*{VRET_R},2)"),
    ("5", "MONTAJE DE ARMADOS Y EQUIPOS", None),
] + [(f"5.{k+1:02d}", f"Montaje de armado tipo {a} – {ARM_DESC[a].lower()}", "Jgo.", rd, {OPE: 1, OFI: 1, PEO: 2 if a in ("PS1-3", "PA1-3") else 3},
      {"Escalera de tres cuerpos": 1, "Camioneta rural 4x4": 0.25} | ({"Camión grúa de 6 t": 0.5} if a == "PSER-3" else {}), [], 0, f"={cA(a)}")
     for k, (a, rd) in enumerate([("PS1-3", 7), ("PA1-3", 5), ("PA2-3", 4), ("TS-3", 4), ("PR3-3", 3), ("DT-3", 2), ("PSER-3", 1), ("PTH-3", 2)])] + [
    ("5.09", "Montaje de reconectador automático, conexionado, puesta en marcha del control", "u", 1.0, {ING: 0.5, OPE: 2, OFI: 1, PEO: 2},
     {"Camión grúa de 6 t": 0.5, "Camioneta rural 4x4": 0.5}, [], 0, f"=PLANILLA!D{PLROW['9.05']}"),
    ("5.10", "Montaje de seccionador fusible tipo cut-out, incluye fusible", "u", 12.0, {OPE: 1, OFI: 1}, {"Escalera de tres cuerpos": 1}, [], 0, f"=PLANILLA!D{PLROW['9.01']}"),
    ("5.11", "Montaje de pararrayos de óxido metálico, incluye conexión a tierra", "u", 12.0, {OPE: 1, OFI: 1}, {"Escalera de tres cuerpos": 1}, [], 0, f"=PLANILLA!D{PLROW['9.04']}"),
    ("6", "MONTAJE DE CONDUCTORES", None),
    ("6.01", "Tendido y puesta en flecha de conductor AAAC 70 mm², por fase", "km", 2.5,
     {OPE: 2, OFI: 3, PEO: 6}, {"Winche de 3 t": 1, "Freno hidráulico": 1, "Caballete porta bobina": 2, "Juego de poleas": 10, "Dinamómetro": 1,
                                "Equipo de comunicación (par)": 2, "Camioneta rural 4x4": 1}, [], 0, f"=ROUND({COND_KM},3)"),
    ("7", "INSTALACIÓN DE PUESTA A TIERRA", None),
    ("7.01", "Excavación en terreno tipo I para pozo de puesta a tierra (Ø 0,80 × 2,60 m)", "m3", 5.0, {OFI: 0.2, PEO: 2}, {}, [], 0,
     f"=ROUND(({cP('PAT-1')}+2*{cP('PAT-2')})*{VPAT},2)"),
    ("7.02", "Instalación de puesta a tierra tipo PAT-1 (una varilla), con cemento conductivo y medición", "u", 5.0, {OPE: 1, PEO: 2}, {"Telurómetro": 0.5, "Camioneta rural 4x4": 0.25}, [], 0, f"={cP('PAT-1')}"),
    ("7.03", "Instalación de puesta a tierra tipo PAT-2 (dos varillas), con cemento conductivo y medición", "u", 3.0, {OPE: 1, PEO: 2}, {"Telurómetro": 1, "Camioneta rural 4x4": 0.25}, [], 0, f"={cP('PAT-2')}"),
    ("7.04", "Relleno y compactación de pozo de puesta a tierra con tierra de cultivo", "m3", 6.0, {OFI: 0.2, PEO: 2}, {}, [(AGU, "m3", 0.05)], 0, "=ROUND(J{r701},2)"),
    ("8", "PRUEBAS Y PUESTA EN SERVICIO", None),
    ("8.01", "Pruebas eléctricas (aislamiento, continuidad, secuencia, PAT) y puesta en servicio de la línea primaria", "km", 5.0,
     {ING: 1, OPE: 1, OFI: 1}, {"Megóhmetro": 1, "Telurómetro": 1, "Camioneta rural 4x4": 1}, [], 0.05, f"=ROUND({Lkm},3)"),
    ("8.02", "Expediente técnico conforme a obra (1 original + 3 copias, digital editable)", "Gbl", 1 / 10, {ING: 1, DIB: 1}, {"Equipo de cómputo y plotter": 1}, [], 0.10, "=1"),
]
APU_TOT = {}; APU_MET = {}; r701 = None
def bp_mo(d): return f"BP!$F${MOROW[d]}"
def bp_eq(d): return f"BP!$F${EQROW[d]}"
for p in PART:
    if p[2] is None: continue
    it, desc, und, rend, mo, eq, mats, pmv, met = p
    put(ap, f"B{r}", f"PARTIDA : {it}", font=FB, border=False); ap.merge_cells(start_row=r, start_column=3, end_row=r, end_column=8)
    put(ap, f"C{r}", desc, font=FB, border=False); r0 = r; r += 1
    put(ap, f"B{r}", "UNIDAD :", border=False); put(ap, f"C{r}", und, border=False)
    put(ap, f"I{r}", "Metrado", font=FI, border=False)
    if "{r701}" in met: met = met.replace("{r701}", str(r701))
    put(ap, f"J{r}", met, font=FB, fmt="#,##0.00"); METC = f"$J${r}"; APU_MET[it] = f"APU!$J${r}"
    if it == "7.01": r701 = r
    r += 1
    put(ap, f"B{r}", "RENDIMIENTO :", border=False); put(ap, f"C{r}", rend, font=FBL, fill=AM, fmt="0.000"); put(ap, f"D{r}", f"{und}/día", border=False)
    REND = f"$C${r}"; r += 1
    fila(ap, r, [None, "DESCRIPCIÓN", "Cuadrilla", "Und.", "Cantidad", "P. Unitario S/", "Parcial S/", "Índice", "", "Cant. total"], font=FB, fill=AZ); r += 1
    # materiales
    put(ap, f"B{r}", "MATERIALES", font=FB); r += 1; m0 = r
    rMOsub = None
    mat_rows = []
    for (md, mu, mq) in mats:
        fila(ap, r, [None, md, None, mu, round(mq, 4), f"=BP!$F${MMROW[md]}", f"=ROUND(E{r}*F{r},2)", "21" if "Cemento" in md else ("38" if md in (HOR, PIE) else "39"),
                     None, f"=E{r}*{METC}"], fmts=[None, None, None, None, "0.0000", MON, MON, None, None, "#,##0.00"]); r += 1
    mv_row = None
    if pmv:
        mv_row = r; fila(ap, r, [None, "Materiales varios", None, "%MO", pmv, None, None, "39", None, None], fmts=[None, None, None, None, "0%", MON, MON]); r += 1
    fila(ap, r, [None, None, None, None, None, "Sub-total", f"=SUM(G{m0}:G{r-1})" if r > m0 else 0], font=FB, fmts=[None] * 6 + [MON]); rM = r; r += 1
    put(ap, f"B{r}", "MANO DE OBRA", font=FB); r += 1; o0 = r
    for d, cq in mo.items():
        fila(ap, r, [None, d, cq, "h-h", f"=C{r}*8/{REND}", f"={bp_mo(d)}", f"=ROUND(E{r}*F{r},2)", "47", None, f"=E{r}*{METC}"],
             fmts=[None, None, "0.00", None, "0.0000", MON, MON, None, None, "#,##0.00"]); r += 1
    fila(ap, r, [None, None, None, None, None, "Sub-total", f"=SUM(G{o0}:G{r-1})"], font=FB, fmts=[None] * 6 + [MON]); rMO = r; r += 1
    if mv_row: ap[f"F{mv_row}"] = f"=G{rMO}"; ap[f"G{mv_row}"] = f"=ROUND(E{mv_row}*F{mv_row},2)"
    put(ap, f"B{r}", "EQUIPOS Y HERRAMIENTAS", font=FB); r += 1; e0 = r
    for d, cq in eq.items():
        fila(ap, r, [None, d, cq, "h-m", f"=C{r}*8/{REND}", f"={bp_eq(d)}", f"=ROUND(E{r}*F{r},2)", "49", None, f"=E{r}*{METC}"],
             fmts=[None, None, "0.00", None, "0.0000", MON, MON, None, None, "#,##0.00"]); r += 1
    fila(ap, r, [None, "Herramientas manuales y EPP (% MO)", None, "%MO", 0.05, f"=G{rMO}", f"=ROUND(E{r}*F{r},2)", "37"], fmts=[None, None, None, None, "0%", MON, MON]); r += 1
    fila(ap, r, [None, None, None, None, None, "Sub-total", f"=SUM(G{e0}:G{r-1})"], font=FB, fmts=[None] * 6 + [MON]); rEQ = r; r += 1
    fila(ap, r, [None, "COSTO UNITARIO DIRECTO", None, None, None, "S/", f"=G{rM}+G{rMO}+G{rEQ}"], font=FB, fill=VE, fmts=[None] * 6 + [MON])
    APU_TOT[it] = f"APU!$G${r}"; r += 3

# =============================================================================== MONTAJE LP
mt = hoja(wb, "MONTAJE LP", [8, 70, 7, 12, 13, 15])
r = encab(mt, "MONTAJE ELECTROMECÁNICO LÍNEA PRIMARIA", "B: MONTAJE ELECTROMECÁNICO DE LÍNEA PRIMARIA", ncol=6)
fila(mt, r, ["ITEM", "DESCRIPCIÓN DE PARTIDAS", "UNID.", "Metrado", "Cost. Unit. S/", "TOTAL S/"], font=FB, fill=AZ); r += 1
MTSUB = {}; grp = None; g0 = None
def cierra(grp, g0):
    global r
    fila(mt, r, [None, f"SUB-TOTAL {grp}", None, None, None, f"=SUM(F{g0}:F{r-1})"], font=FB, fmts=[None] * 5 + [MON]); MTSUB[grp] = r; r += 2
MTROW = {}
for p in PART:
    if p[2] is None:
        if grp: cierra(grp, g0)
        grp = p[0]; put(mt, f"A{r}", p[0], font=FB); put(mt, f"B{r}", p[1], font=FB); r += 1; g0 = r; continue
    it, desc, und = p[0], p[1], p[2]
    fila(mt, r, [it, desc, und, f"={APU_MET[it]}", f"={APU_TOT[it]}", f"=ROUND(D{r}*E{r},2)"], fmts=[None, None, None, "#,##0.00", MON, MON])
    mt[f"B{r}"].alignment = LEFT; MTROW[it] = r; r += 1
cierra(grp, g0)
fila(mt, r, [None, "TOTAL MONTAJE ELECTROMECÁNICO", None, None, None, "=" + "+".join(f"F{v}" for v in MTSUB.values())], font=FB, fill=VE, fmts=[None] * 5 + [MON])
MT_TOT = f"'MONTAJE LP'!F{r}"
nota(mt, r + 2, "Nota: el contratista coordinará con la empresa concesionaria la codificación y señalización final de los postes. Metrados sustentados "
     "en la PLANILLA y en los Anexos 5 (cimentaciones) y 7 (retenidas) del Volumen VI.", ncol=6, h=30)

# =============================================================================== TRANSPORTE LP
tr = hoja(wb, "TRANSPORTE LP", [8, 66, 9, 15, 15])
r = encab(tr, f"TRANSPORTE DE MATERIALES LÍNEA PRIMARIA (proveedor → almacén de obra en {FIN['nombre']})", "C: TRANSPORTE DE MATERIALES", ncol=5)
put(tr, f"B{r}", "Relación transporte / suministro adoptada (proyecto.yaml → economico.transporte_pct; referencia: proyecto Ayaviri–Pucará 7,0 %)", font=FB)
put(tr, f"D{r}", TRANSPORTE_PCT, font=FBL, fill=AM, fmt="0.00%"); PTR = f"$D${r}"; r += 2
fila(tr, r, ["ITEM", "GRUPO DE SUMINISTRO", "% ", "Suministro S/", "Transporte S/"], font=FB, fill=AZ); r += 1; t0 = r
for g, gname in GRUPOS:
    fila(tr, r, [g, gname, f"={PTR}", f"='SUMINISTRO LP'!F{SUSUB[g]}", f"=ROUND(C{r}*D{r},2)"], fmts=[None, None, "0.00%", MON, MON]); r += 1
fila(tr, r, [None, "TOTAL TRANSPORTE DE MATERIALES", None, f"=SUM(D{t0}:D{r-1})", f"=SUM(E{t0}:E{r-1})"], font=FB, fill=VE, fmts=[None, None, None, MON, MON])
TR_TOT = f"'TRANSPORTE LP'!E{r}"
nota(tr, r + 2, "SUPUESTO: a falta de cotización de fletes, se aplica la relación del expediente de referencia LP Ayaviri–Pucará. El postor sustentará su costo de "
     f"transporte ({ECON.get('ruta_transporte', 'ruta de transporte a sustentar')}).", ncol=5, h=30)

# =============================================================================== GG
gg = hoja(wb, "GG", [52, 8, 7, 9, 9, 12, 14])
r = encab(gg, "CÁLCULO DE LOS GASTOS GENERALES DIRECTOS E INDIRECTOS", ncol=7)
put(gg, f"A{r}", "Plazo de obra (días calendario)", font=FB); put(gg, f"E{r}", PLAZO, font=FBL, fill=AM); PZ = f"GG!$E${r}"
put(gg, f"F{r}", "C.D. S/", font=FB); put(gg, f"G{r}", "='RESUMEN OBRA'!F14", fmt=MON); CD = f"GG!$G${r}"; r += 2
def bloque(titulo, items, kind="dia"):
    """items: (desc, und, cant, part, tiempo|None→plazo, costo/día)."""
    global r
    titulo_seccion(gg, r, titulo, ncol=7); r += 1
    fila(gg, r, ["Descripción", "Und.", "Cant.", "Particip.", "Tiempo (días)", "Costo/día S/", "Parcial S/"], font=FB, fill=AZ); r += 1; b0 = r
    for d, u, c, pa, t, cd in items:
        fila(gg, r, [d, u, c, pa, t if t is not None else f"={PZ}", cd, f"=ROUND(C{r}*D{r}*E{r}*F{r},2)"], fmts=[None, None, "0", "0%", "0", MON, MON])
        gg[f"F{r}"].fill = AM; gg[f"F{r}"].font = FBL; r += 1
    fila(gg, r, ["PARCIAL", None, None, None, None, None, f"=SUM(G{b0}:G{r-1})"], font=FB, fmts=[None] * 6 + [MON]); rr = r; r += 2
    return rr
titulo_seccion(gg, r, "1.0  GASTOS GENERALES VARIABLES – DIRECTOS (relacionados con el plazo)", ncol=7, fill=VE); r += 1
p11 = bloque("1.1 Personal profesional principal", [("Ingeniero residente de obra (colegiado y habilitado)", "Und.", 1, 1, None, 250),
                                                    ("Ingeniero de seguridad y salud en el trabajo", "Und.", 1, 1, None, 220)])
p12 = bloque("1.2 Personal auxiliar y de apoyo", [("Administrador", "Und.", 1, 1, None, 120), ("Dibujante CAD", "Und.", 1, 0.5, None, 90), ("Almacenero", "Und.", 1, 0.9, None, 70),
                                                  ("Chofer", "Und.", 1, 1, None, 70), ("Guardián", "Und.", 1, 1, None, 70)])
p13 = bloque("1.3 Hospedaje, oficina, almacén y comunicaciones", [("Hospedaje", "Glb/día", 1, 1, None, 15), ("Alquiler de oficina en obra (incluye agua y luz)", "Glb/día", 1, 1, None, 15),
                                                                  ("Alquiler de almacén (incluye agua y luz)", "Glb/día", 1, 1, None, 15), ("Comunicaciones: teléfono, internet, radio", "Glb/día", 1, 1, None, 18),
                                                                  ("Alimentación", "Glb/día", 1, 1, None, 70), ("Pasajes", "Glb/día", 1, 1, None, 20)])
p14 = bloque("1.4 Mobiliario, equipo, movilidad y útiles", [("Mobiliario de oficina (desgaste)", "Glb/día", 2, 0.072, None, 40), ("Computadora (desgaste)", "Glb/día", 2, 0.06, None, 35),
                                                            ("Impresora (desgaste)", "Glb/día", 1, 0.06, None, 10), ("Útiles de oficina", "Glb/día", 1, 1, None, 10),
                                                            ("Camioneta pick-up 4x4 (alquiler)", "Glb/día", 1, 1, None, 200), ("Combustible", "Glb/día", 1, 1, None, 100),
                                                            ("Mantenimiento", "Glb/día", 1, 1, None, 30)])
titulo_seccion(gg, r, "1.5 Gastos financieros y seguros", ncol=7); r += 1
fila(gg, r, ["Descripción", "Und.", "% CD", "Tasa anual", "Tiempo (días)", "Costo/día S/", "Parcial S/"], font=FB, fill=AZ); r += 1; b0 = r
fila(gg, r, ["Carta fianza de fiel cumplimiento (10 % del contrato)", "día", 0.10, 0.02, f"={PZ}+90", f"={CD}*C{r}*D{r}/360", f"=ROUND(E{r}*F{r},2)"], fmts=[None, None, "0%", "0.0%", "0", MON, MON]); r += 1
fila(gg, r, ["Seguros: CAR, responsabilidad civil y SCTR administrativo", "día", 0.005, None, f"={PZ}", f"={CD}*C{r}/{PZ}", f"=ROUND(E{r}*F{r},2)"], fmts=[None, None, "0.0%", None, "0", MON, MON]); r += 1
fila(gg, r, ["PARCIAL", None, None, None, None, None, f"=SUM(G{b0}:G{r-1})"], font=FB, fmts=[None] * 6 + [MON]); p15 = r; r += 2
fila(gg, r, ["TOTAL GASTOS GENERALES VARIABLES (1)", None, None, None, None, None, f"=G{p11}+G{p12}+G{p13}+G{p14}+G{p15}"], font=FB, fill=VE, fmts=[None] * 6 + [MON]); GGV = r; r += 2
titulo_seccion(gg, r, "2.0  GASTOS GENERALES FIJOS – INDIRECTOS", ncol=7, fill=VE); r += 1
p21 = bloque("2.1 Personal de la sede central (participación parcial)", [("Gerente", "Und.", 1, 0.2, None, 200), ("Contador", "Und.", 1, 0.2, None, 150),
                                                                        ("Secretaria", "Und.", 1, 0.2, None, 80)])
p22 = bloque("2.2 Gastos de oficina principal", [("Oficina (incluye agua y luz)", "Und.", 1, 0.2, None, 60), ("Comunicaciones", "Glb", 1, 0.2, None, 40),
                                                 ("Material y equipos de oficina", "Glb", 1, 0.2, None, 50)])
p23 = bloque("2.3 Operación experimental asistida (30 días)", [("Ingeniero residente", "Und.", 1, 1, 30, 316.67), ("Operario", "Und.", 1, 1, 30, 156.67),
                                                              ("Camioneta pick-up con combustible", "Glb/día", 1, 1, 30, 73.0)])
titulo_seccion(gg, r, "2.4 Preparación de oferta e imprevistos", ncol=7); r += 1
fila(gg, r, ["Elaboración de propuesta (0,03 % del CD)", "Gbl", 1, 1, 1, f"={CD}*0.0003", f"=ROUND(F{r},2)"], fmts=[None, None, "0", "0%", "0", MON, MON]); r += 1
fila(gg, r, ["Imprevistos (0,03 % del CD)", "Gbl", 1, 1, 1, f"={CD}*0.0003", f"=ROUND(F{r},2)"], fmts=[None, None, "0", "0%", "0", MON, MON]); r += 1
fila(gg, r, ["PARCIAL", None, None, None, None, None, f"=SUM(G{r-2}:G{r-1})"], font=FB, fmts=[None] * 6 + [MON]); p24 = r; r += 2
fila(gg, r, ["TOTAL GASTOS GENERALES FIJOS (2)", None, None, None, None, None, f"=G{p21}+G{p22}+G{p23}+G{p24}"], font=FB, fill=VE, fmts=[None] * 6 + [MON]); GGF = r; r += 2
titulo_seccion(gg, r, "3.0  RESUMEN", ncol=7); r += 1
fila(gg, r, ["Gastos generales variables – directos", None, None, None, None, f"=G{GGV}/{CD}", f"=G{GGV}"], fmts=[None] * 5 + ["0.0000%", MON]); GGVP = f"GG!$F${r}"; r += 1
fila(gg, r, ["Gastos generales fijos – indirectos", None, None, None, None, f"=G{GGF}/{CD}", f"=G{GGF}"], fmts=[None] * 5 + ["0.0000%", MON]); GGFP = f"GG!$F${r}"; r += 1
fila(gg, r, ["TOTAL GASTOS GENERALES", None, None, None, None, f"=F{r-2}+F{r-1}", f"=G{r-2}+G{r-1}"], font=FB, fmts=[None] * 5 + ["0.0000%", MON]); r += 2
nota(gg, r, "Costos diarios referenciales tomados del expediente de referencia LP Ayaviri–Pucará (2022); actualizar a la fecha de convocatoria. Plazo de ejecución: " + str(PLAZO) + " días calendario.", ncol=7)

# =============================================================================== GS (supervisión)
gs = hoja(wb, "GS", [52, 8, 7, 9, 9, 12, 14])
r = encab(gs, "ANÁLISIS DE GASTOS DE SUPERVISIÓN", ncol=7)
put(gs, f"A{r}", "Plazo de obra (días)", font=FB); put(gs, f"E{r}", f"={PZ}"); r += 2
PZ2 = f"GS!$E${r-2}"
def bloque_s(titulo, items):
    global r
    titulo_seccion(gs, r, titulo, ncol=7); r += 1
    fila(gs, r, ["Descripción", "Und.", "Cant.", "Particip.", "Tiempo (días)", "Costo/día S/", "Parcial S/"], font=FB, fill=AZ); r += 1; b0 = r
    for d, u, c, pa, t, cd in items:
        fila(gs, r, [d, u, c, pa, t if t is not None else f"={PZ2}", cd, f"=ROUND(C{r}*D{r}*E{r}*F{r},2)"], fmts=[None, None, "0", "0%", "0", MON, MON])
        gs[f"F{r}"].fill = AM; gs[f"F{r}"].font = FBL; r += 1
    fila(gs, r, ["PARCIAL", None, None, None, None, None, f"=SUM(G{b0}:G{r-1})"], font=FB, fmts=[None] * 6 + [MON]); rr = r; r += 2; return rr
s1 = bloque_s("1.1 Personal profesional y técnico", [("Supervisor de obra (ingeniero electricista o mecánico electricista)", "Und.", 1, 1, None, 250), ("Técnico conductor", "Und.", 1, 1, None, 75)])
s2 = bloque_s("1.2 Oficina, hospedaje y comunicaciones", [("Alquiler de oficina en obra", "Glb", 1, 1, None, 20), ("Comunicaciones", "Glb", 1, 1, None, 10), ("Hospedaje y alimentación", "Glb", 1, 1, None, 50)])
s3 = bloque_s("1.3 Mobiliario, equipo y movilidad", [("Equipo de cómputo, GPS y cámara (desgaste)", "Glb", 1, 0.1, None, 120), ("Útiles de oficina", "Glb", 1, 0.5, None, 15),
                                                       ("Camioneta 4x4 doble cabina", "Glb", 1, 1, None, 150), ("Combustible y mantenimiento", "Glb", 1, 1, None, 80)])
s4 = bloque_s("2.0 Liquidación de obra – supervisión", [("Liquidación técnica y financiera de obra", "Glb", 1, 1, 1, 7000)])
fila(gs, r, ["SUB TOTAL", None, None, None, None, None, f"=G{s1}+G{s2}+G{s3}+G{s4}"], font=FB, fmts=[None] * 6 + [MON]); st = r; r += 1
fila(gs, r, ["Utilidad (10 %)", None, None, None, None, 0.10, f"=ROUND(G{st}*F{r},2)"], fmts=[None] * 5 + ["0%", MON]); r += 1
fila(gs, r, ["IGV (18 %)", None, None, None, None, IGV_PCT, f"=ROUND((G{st}+G{r-1})*F{r},2)"], fmts=[None] * 5 + ["0%", MON]); r += 1
fila(gs, r, ["TOTAL SUPERVISIÓN DE OBRA (incluye IGV)", None, None, None, None, None, f"=G{st}+G{r-2}+G{r-1}"], font=FB, fill=VE, fmts=[None] * 6 + [MON]); GS_TOT = f"GS!$G${r}"

# =============================================================================== CCL
cc = hoja(wb, "CCL", [52, 8, 7, 9, 9, 12, 14])
r = encab(cc, "GASTOS DE COORDINACIÓN Y DE LIQUIDACIÓN FINANCIERA DE OBRA (a cargo de la Entidad)", ncol=7)
titulo_seccion(cc, r, "1. Gastos de coordinación de obra", ncol=7); r += 1
fila(cc, r, ["Descripción", "Und.", "Cant.", "Particip.", "Tiempo (días)", "Costo/día S/", "Parcial S/"], font=FB, fill=AZ); r += 1
fila(cc, r, ["Coordinador de obra de la Entidad", "Und.", 1, 1, f"={PZ}", 250, f"=ROUND(C{r}*D{r}*E{r}*F{r},2)"], fmts=[None, None, "0", "0%", "0", MON, MON]); r += 1
fila(cc, r, ["IGV (18 %)", None, None, None, None, IGV_PCT, f"=ROUND(G{r-1}*F{r},2)"], fmts=[None] * 5 + ["0%", MON]); r += 1
fila(cc, r, ["PARCIAL COORDINACIÓN", None, None, None, None, None, f"=G{r-2}+G{r-1}"], font=FB, fmts=[None] * 6 + [MON]); CC1 = f"CCL!$G${r}"; r += 2
titulo_seccion(cc, r, "2. Gastos de liquidación financiera de obra", ncol=7); r += 1
fila(cc, r, ["Liquidador de obra", "Glb/día", 1, 1, 30, 266.67, f"=ROUND(C{r}*D{r}*E{r}*F{r},2)"], fmts=[None, None, "0", "0%", "0", MON, MON]); r += 1
fila(cc, r, ["IGV (18 %)", None, None, None, None, IGV_PCT, f"=ROUND(G{r-1}*F{r},2)"], fmts=[None] * 5 + ["0%", MON]); r += 1
fila(cc, r, ["PARCIAL LIQUIDACIÓN", None, None, None, None, None, f"=G{r-2}+G{r-1}"], font=FB, fmts=[None] * 6 + [MON]); CC2 = f"CCL!$G${r}"

# =============================================================================== RESUMEN OBRA
ro = hoja(wb, "RESUMEN OBRA", [14, 60, 10, 8, 4, 17])
ro.merge_cells("A2:F2"); ro["A2"] = "RESUMEN GENERAL DEL VALOR REFERENCIAL DE OBRA"; ro["A2"].font = FT; ro["A2"].alignment = CEN
put(ro, "A4", "PROYECTO", font=FB, border=False); ro.merge_cells("B4:F4"); put(ro, "B4", NOMBRE, border=False)
put(ro, "A5", "UBICACIÓN", font=FB, border=False); ro.merge_cells("B5:F5"); put(ro, "B5", UBIC, border=False)
put(ro, "A6", "PLAZO", font=FB, border=False); put(ro, "B6", f"={PZ}&\" días calendario\"", border=False)
put(ro, "A7", "FECHA", font=FB, border=False); put(ro, "B7", f"{PROYECTO['fecha']} (precios sin IGV a la fecha indicada en BP)", border=False)
put(ro, "A8", "PROYECTISTA", font=FB, border=False); put(ro, "B8", PROY, border=False)
fila(ro, 10, ["ITEM", "DESCRIPCIÓN", "%", None, None, "TOTAL S/"], font=FB, fill=AZ)
rows = [("A", "SUMINISTRO DE MATERIALES", None, f"=ROUND({SU_TOT},2)"), ("B", "MONTAJE ELECTROMECÁNICO", None, f"=ROUND({MT_TOT},2)"),
        ("C", "TRANSPORTE DE MATERIALES", None, f"=ROUND({TR_TOT},2)")]
r = 11
for a, d, p, f in rows:
    fila(ro, r, [a, d, None, None, None, f], fmts=[None] * 5 + [MON]); r += 1
r += 0
fila(ro, 14, ["D", "COSTO DIRECTO (C.D.)", None, None, None, "=SUM(F11:F13)"], font=FB, fmts=[None] * 5 + [MON])
fila(ro, 15, ["E", "GASTOS GENERALES DIRECTOS", f"={GGVP}", "de D", None, "=ROUND(F14*C15,2)"], fmts=[None, None, "0.0000%", None, None, MON])
fila(ro, 16, ["F", "GASTOS GENERALES INDIRECTOS", f"={GGFP}", "de D", None, "=ROUND(F14*C16,2)"], fmts=[None, None, "0.0000%", None, None, MON])
fila(ro, 17, ["G", "UTILIDAD", UTILIDAD_PCT, "de D", None, "=ROUND(F14*C17,2)"], fmts=[None, None, "0.00%", None, None, MON])
ro["C17"].fill = AM; ro["C17"].font = FBL
fila(ro, 18, ["H", "SUB-TOTAL SIN IGV (S/)", None, None, None, "=F14+F15+F16+F17"], font=FB, fmts=[None] * 5 + [MON])
fila(ro, 19, ["I", "IMPUESTO GENERAL A LAS VENTAS IGV (18 %)", IGV_PCT, None, None, "=ROUND(F18*C19,2)"], fmts=[None, None, "0%", None, None, MON])
fila(ro, 20, ["J", "VALOR REFERENCIAL DE OBRA (incluye IGV) S/", None, None, None, "=F18+F19"], font=FB, fill=VE, fmts=[None] * 5 + [MON])
fila(ro, 22, ["K", "GASTOS DE SUPERVISIÓN DE OBRA (incluye IGV)", None, None, None, f"={GS_TOT}"], fmts=[None] * 5 + [MON])
fila(ro, 23, ["L", "GASTOS DE COORDINACIÓN DE OBRA (incluye IGV)", None, None, None, f"={CC1}"], fmts=[None] * 5 + [MON])
fila(ro, 24, ["M", "GASTOS DE LIQUIDACIÓN FINANCIERA DE OBRA (incluye IGV)", None, None, None, f"={CC2}"], fmts=[None] * 5 + [MON])
fila(ro, 25, ["N", "ELABORACIÓN DEL EXPEDIENTE TÉCNICO, DIA Y CIRA", None, None, None, "No incluido"], fmts=[None] * 6)
fila(ro, 27, ["O", "COSTO TOTAL DE INVERSIÓN (J + K + L + M) S/", None, None, None, "=F20+F22+F23+F24"], font=FB, fill=VE, fmts=[None] * 5 + [MON])
nota(ro, 29, "Notas: (1) Los costos de expediente técnico, DIA, CIRA/PMA y servidumbre no se incluyen por no contar con contratos o "
     "cotizaciones; la Entidad los agregará. (2) Precios de suministro con fuente indicada en la hoja BP (cotizaciones del expediente de referencia convertidas "
     "al tipo de cambio vigente y estimados): deben actualizarse con un mínimo de tres cotizaciones antes de la convocatoria.", ncol=6, h=55)

# =============================================================================== CRONOGRAMAS (Gantt, valorizado, suministro)
_f = PLAZO / 120.0
def _d(x): return max(1, min(PLAZO, int(round(x * _f))))
ACT = [(it, d, (1 if a == 1 else _d(a)), _d(b), g) for it, d, a, b, g in
       [("1", "OBRAS PRELIMINARES", 1, 120, "1"), ("2", "TRANSPORTE DE MATERIALES DE ALMACÉN A PUNTO DE INSTALACIÓN", 31, 95, "2"),
        ("3", "INSTALACIÓN DE POSTES", 31, 85, "3"), ("4", "INSTALACIÓN DE RETENIDAS", 51, 95, "4"), ("5", "MONTAJE DE ARMADOS Y EQUIPOS", 46, 100, "5"),
        ("6", "MONTAJE DE CONDUCTORES", 61, 105, "6"), ("7", "INSTALACIÓN DE PUESTA A TIERRA", 41, 95, "7"), ("8", "PRUEBAS Y PUESTA EN SERVICIO", 106, 120, "8")]]
NSEM = math.ceil(PLAZO / 7)
gt = hoja(wb, "GANTT", [6, 52, 8, 8, 8, 14] + [3.2] * NSEM, apaisado=True)
r = encab(gt, f"CRONOGRAMA DE EJECUCIÓN DE OBRA (DIAGRAMA GANTT) – {PLAZO} DÍAS CALENDARIO", ncol=6 + NSEM)
fila(gt, r, ["ITEM", "ACTIVIDAD", "Inicio (día)", "Fin (día)", "Duración", "Monto S/"] + [f"S{k}" for k in range(1, NSEM + 1)], font=FB, fill=AZ); r += 1
gt0 = r
fila(gt, r, ["A", "SUMINISTRO Y TRANSPORTE DE MATERIALES A ALMACÉN DE OBRA", 1, _d(60), "=D{0}-C{0}+1".format(r), f"=ROUND({SU_TOT}+{TR_TOT},2)"], fmts=[None] * 5 + [MON]); r += 1
for it, d, a, b, g in ACT:
    fila(gt, r, [it, d, a, b, f"=D{r}-C{r}+1", f"='MONTAJE LP'!F{MTSUB[g]}"], fmts=[None] * 5 + [MON]); r += 1
gt1 = r - 1
from openpyxl.formatting.rule import FormulaRule
rng = f"G{gt0}:{CL(6 + NSEM)}{gt1}"
for rr in range(gt0, gt1 + 1):
    for k in range(NSEM):
        c = gt.cell(rr, 7 + k); c.border = BOX
        c.value = f"=IF(AND({7*k+7}>=$C{rr},{7*k+1}<=$D{rr}),1,0)"; c.font = Font(name="Arial", size=6, color="FFFFFF")
gt.conditional_formatting.add(rng, FormulaRule(formula=[f"G{gt0}=1"], fill=PatternFill("solid", fgColor="2F5597"), font=Font(color="2F5597")))
nota(gt, r + 1, f"S1…S{NSEM}: semanas de obra. Actividades de montaje con montos enlazados a los subtotales de la hoja MONTAJE LP. "
     "Archivo CRONOGRAMA.xml (importable en MS Project) con las mismas actividades.", ncol=min(24, 6 + NSEM), h=30)
NMES = math.ceil(PLAZO / 30)
meses = [(30 * k + 1, min(PLAZO, 30 * (k + 1))) for k in range(NMES)]
MC = [CL(4 + k) for k in range(NMES)]; CTRL = CL(4 + NMES)
cv = hoja(wb, "CRONOG_VALORIZ", [6, 52, 16] + [15] * NMES + [12])
r = encab(cv, "CALENDARIO VALORIZADO DE AVANCE DE OBRA", ncol=4 + NMES)
fila(cv, r, ["ITEM", "DESCRIPCIÓN", "TOTAL S/"] + [f"MES {k+1}" for k in range(NMES)] + ["Control"], font=FB, fill=AZ); r += 1; cv0 = r
for it, ds, ref in (("A", "SUMINISTRO DE MATERIALES", SU_TOT), ("C", "TRANSPORTE DE MATERIALES", TR_TOT)):
    v = [it, ds, f"={ref}", f"=ROUND(C{r}*0.3,2)"] + ([f"=C{r}-D{r}"] if NMES > 1 else []) + [0] * max(0, NMES - 2)
    if NMES == 1: v[3] = f"=C{r}"
    fila(cv, r, v + [f"=SUM(D{r}:{MC[-1]}{r})-C{r}"], fmts=[None, None] + [MON] * (NMES + 2)); r += 1
for it, d, a, b, g in ACT:
    gr = gt0 + 1 + [x[0] for x in ACT].index(it)
    vals = [f"B.{it}", d, f"=GANTT!F{gr}"]
    for k, (m0, m1) in enumerate(meses):
        vals.append(f"=ROUND(C{r}*MAX(0,MIN(GANTT!D{gr},{m1})-MAX(GANTT!C{gr},{m0})+1)/GANTT!E{gr},2)" if k < NMES - 1 else
                    (f"=C{r}-SUM(D{r}:{MC[-2]}{r})" if NMES > 1 else f"=C{r}"))
    vals.append(f"=SUM(D{r}:{MC[-1]}{r})-C{r}")
    fila(cv, r, vals, fmts=[None, None] + [MON] * (NMES + 2)); r += 1
fila(cv, r, [None, "COSTO DIRECTO", f"=SUM(C{cv0}:C{r-1})"] + [f"=SUM({c}{cv0}:{c}{r-1})" for c in MC] + [None], font=FB, fmts=[None, None] + [MON] * (NMES + 1)); rcd = r; r += 1
for lab, ref in [("GASTOS GENERALES DIRECTOS", "'RESUMEN OBRA'!$C$15"), ("GASTOS GENERALES INDIRECTOS", "'RESUMEN OBRA'!$C$16"), ("UTILIDAD", "'RESUMEN OBRA'!$C$17")]:
    fila(cv, r, [None, lab, f"=ROUND(C{rcd}*{ref},2)"] + [f"=ROUND({c}{rcd}*{ref},2)" for c in MC] + [None], fmts=[None, None] + [MON] * (NMES + 1)); r += 1
fila(cv, r, [None, "SUB-TOTAL", f"=SUM(C{rcd}:C{r-1})"] + [f"=SUM({c}{rcd}:{c}{r-1})" for c in MC] + [None], font=FB, fmts=[None, None] + [MON] * (NMES + 1)); rst = r; r += 1
fila(cv, r, [None, f"IGV ({IGV_PCT*100:.0f} %)", f"=ROUND(C{rst}*{IGV_PCT},2)"] + [f"=ROUND({c}{rst}*{IGV_PCT},2)" for c in MC] + [None], fmts=[None, None] + [MON] * (NMES + 1)); r += 1
fila(cv, r, [None, "TOTAL MENSUAL", f"=C{rst}+C{r-1}"] + [f"={c}{rst}+{c}{r-1}" for c in MC] + [None], font=FB, fill=VE, fmts=[None, None] + [MON] * (NMES + 1)); rtt = r; r += 1
fila(cv, r, [None, "% DE AVANCE MENSUAL", None] + [f"={c}{rtt}/$C${rtt}" for c in MC] + [None], fmts=[None, None, None] + ["0.00%"] * NMES); r += 1
fila(cv, r, [None, "% DE AVANCE ACUMULADO", None] + [f"=D{r-1}"] + [f"={MC[k-1]}{r}+{MC[k]}{r-1}" for k in range(1, NMES)] + [None], fmts=[None, None, None] + ["0.00%"] * NMES)
cs = hoja(wb, "CRONOG_SUMINISTRO", [6, 52, 16, 15, 15, 15])
r = encab(cs, "CRONOGRAMA DE SUMINISTRO VALORIZADO (entregas en almacén de obra)", ncol=6)
fila(cs, r, ["ITEM", "GRUPO", "TOTAL S/", "MES 1 (30 %)", "MES 2 (70 %)", "Plazo de entrega"], font=FB, fill=AZ); r += 1; c0_ = r
plazos = {"1": "45 días", "2": "45 días", "3": "60 días", "4": "45 días", "5": "30 días", "6": "45 días", "7": "45 días", "8": "30 días", "9": "60 días (reconectador 90 días)", "10": "30 días"}
for g, gname in GRUPOS:
    fila(cs, r, [g, gname, f"='SUMINISTRO LP'!F{SUSUB[g]}", f"=ROUND(C{r}*0.3,2)", f"=C{r}-D{r}", plazos[g]], fmts=[None, None, MON, MON, MON, None]); r += 1
fila(cs, r, [None, "TOTAL", f"=SUM(C{c0_}:C{r-1})", f"=SUM(D{c0_}:D{r-1})", f"=SUM(E{c0_}:E{r-1})", None], font=FB, fill=VE, fmts=[None, None, MON, MON, MON])

# =============================================================================== FP
fp = hoja(wb, "FP", [6, 46, 9, 10, 15, 11, 11, 11])
r = encab(fp, "FÓRMULA POLINÓMICA DE REAJUSTE", ncol=8)
put(fp, f"B{r}", "PRESUPUESTO BASE (sin IGV) S/", font=FB); put(fp, f"E{r}", "='RESUMEN OBRA'!F18", fmt=MON); PB = f"$E${r}"; r += 2
fila(fp, r, ["N.°", "ELEMENTO REPRESENTATIVO", "ÍNDICE", "SÍMBOLO", "MONTO S/", "INCIDENCIA", "% PARCIAL", "COEFICIENTE"], font=FB, fill=AZ); r += 1
def gs_(g): return f"'SUMINISTRO LP'!F{SUSUB[g]}"
GGU = "('RESUMEN OBRA'!F15+'RESUMEN OBRA'!F16+'RESUMEN OBRA'!F17)"
mono = [("1", [("POSTES DE CONCRETO", "62", "PC", gs_("1"))]),
        ("2", [("CONDUCTOR AAAC Y ACCESORIOS", "30", "CO", f"{gs_('3')}+{gs_('4')}"), ("EQUIPOS DE PROTECCIÓN Y MANIOBRA", "06", "PR", gs_("9"))]),
        ("3", [("AISLADORES Y ACCESORIOS", "11", "AI", gs_("2")), ("FERRETERÍA, RETENIDAS Y SEÑALIZACIÓN", "02", "FE", f"{gs_('6')}+{gs_('7')}+{gs_('10')}"),
               ("CONDUCTOR DE COBRE Y PUESTA A TIERRA", "06", "CC", f"{gs_('5')}+{gs_('8')}")]),
        ("4", [("MONTAJE ELECTROMECÁNICO", "47", "J", MT_TOT), ("TRANSPORTE", "32", "T", TR_TOT)]),
        ("5", [("GASTOS GENERALES Y UTILIDAD", "39", "GU", GGU)])]
f0 = r; firsts = []
for n_, parts in mono:
    m_first = r; firsts.append(r)
    for k, (d, ind, sym, ref) in enumerate(parts):
        fila(fp, r, [n_ if k == 0 else None, d, ind, sym, f"={ref}", f"=E{r}/{PB}", None, None], fmts=[None, None, None, None, MON, "0.000", "0.000", "0.000"]); r += 1
    for rr in range(m_first, r):
        fp[f"G{rr}"] = f"=E{rr}/SUM(E{m_first}:E{r-1})"; fp[f"G{rr}"].number_format = "0.000"
    fp[f"H{m_first}"] = (f"=ROUND(SUM(F{m_first}:F{r-1}),3)" if n_ != "5" else "=1-" + "-".join(f"H{x}" for x in firsts[:-1]))
    fp[f"H{m_first}"].number_format = "0.000"
fila(fp, r, [None, "TOTAL", None, None, f"=SUM(E{f0}:E{r-1})", f"=SUM(F{f0}:F{r-1})", None, f"=SUM(H{f0}:H{r-1})"], font=FB, fmts=[None, None, None, None, MON, "0.000", None, "0.000"]); rT = r; r += 2
kf = "=\"K = \"" + "".join(f"&TEXT(H{fr},\"0.000\")&\"·{s}r/{s}o\"" + ("&\" + \"" if i < len(firsts) - 1 else "") for i, (fr, s) in enumerate(zip(firsts, ["PC", "CO", "AI", "J", "GU"])))
fp.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8); put(fp, f"B{r}", kf, font=FB); r += 1
put(fp, f"B{r}", f'=IF(ABS(H{rT}-1)<0.0005,"Suma de coeficientes = 1,000: CONFORME","AJUSTAR REDONDEO DEL ÚLTIMO MONOMIO")', font=FI); r += 2
nota(fp, r, "Notas: (1) Subíndice 'o' = índice a la fecha del presupuesto base; 'r' = índice a la fecha de valorización (D.S. 011-79-VC y "
     "modificatorias). (2) En monomios compuestos se adopta el índice del elemento de mayor incidencia. (3) Códigos de Índices Unificados de "
     "Precios (INEI) iguales a los del expediente de referencia: verificar contra la tabla vigente antes de la convocatoria. (4) Ningún monomio menor al 5 %.", ncol=8, h=55)

# =============================================================================== Relación de insumos (MO, EQUIPO, MATERIALES)
for nombre, dic, unidad in (("MO", MO, "h-h"), ("EQUIPO", EQ, "h-m"), ("MATERIALES", MATMON, None)):
    ws = hoja(wb, nombre, [62, 8, 14, 13, 15])
    r = encab(ws, f"RELACIÓN DE INSUMOS – {'MANO DE OBRA' if nombre == 'MO' else ('EQUIPOS' if nombre == 'EQUIPO' else 'MATERIALES DE MONTAJE')}", ncol=5)
    fila(ws, r, ["Descripción", "Und.", "Cantidad total", "Precio S/", "Parcial S/"], font=FB, fill=AZ); r += 1; i0 = r
    for d in dic:
        u = unidad or dic[d][0]
        pref = f"BP!$F${MOROW[d]}" if nombre == "MO" else (f"BP!$F${EQROW[d]}" if nombre == "EQUIPO" else f"BP!$F${MMROW[d]}")
        fila(ws, r, [d, u, f'=SUMIF(APU!$B:$B,A{r},APU!$J:$J)', f"={pref}", f"=ROUND(C{r}*D{r},2)"], fmts=[None, None, "#,##0.00", MON, MON]); r += 1
    fila(ws, r, ["TOTAL", None, None, None, f"=SUM(E{i0}:E{r-1})"], font=FB, fmts=[None] * 4 + [MON])

order = ["RESUMEN OBRA", "GG", "GS", "CCL", "CRONOG_VALORIZ", "CRONOG_SUMINISTRO", "GANTT", "SUMINISTRO LP", "MONTAJE LP", "TRANSPORTE LP",
         "APU", "FP", "PLANILLA", "BOM_ARMADOS", "BP", "MO", "EQUIPO", "MATERIALES"]
wb._sheets = [wb[n] for n in order]
wb.save(OUT)
json.dump({"ACT": ACT, "SUSUB": SUSUB, "MTSUB": MTSUB, "PLROW": PLROW}, open(W / "presup_meta.json", "w"))
print("ok presupuesto", OUT)
