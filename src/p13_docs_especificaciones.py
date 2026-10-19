"""Volumen IV: Especificaciones técnicas de suministro (ETS-LP-xx con tablas de datos técnicos garantizados) y de montaje
electromecánico (generales + particulares por partida del presupuesto). Uso: python p13_docs_especificaciones.py WORK OUTDIR"""
import json, sys
from pathlib import Path
import openpyxl
sys.path.insert(0, str(Path(__file__).parent))
from params import *
from docgen import *
from comun import cargar, lista, fusibles_txt

W = Path(sys.argv[1]); OUT = Path(sys.argv[2])
C = cargar(W); K = C["K"]; DER = C["ders"]; ALT = C["ALT"]; CONC = PROYECTO["concesionaria"]
_n = lambda x, d=0: f"{float(x):,.{d}f}".replace(",", " ").replace(".", ",").replace("-", "−")
T_AMB_MAX = CLIMA.get("t_max_ambiente", round(float(CLIMA.get("t_max_media", 15)) + 10))
COND_OP = [("Tensión nominal / máxima del sistema", "kV", "22,9 / 25"), ("Frecuencia", "Hz", "60"), ("Sistema", "–", "Trifásico, neutro efectivamente a tierra"),
           ("Altitud de instalación", "m s.n.m.", f"hasta {_n(ALT)} (diseño)"), ("Temperatura ambiente mínima / media / máxima", "°C", f"{_n(CLIMA['t_min_diseno'])} / {_n(T_MEDIA)} / {_n(T_AMB_MAX)}"),
           ("Humedad relativa", "%", "30 a 95"), ("Nivel de contaminación (IEC 60815)", "–", f"{CLIMA.get('contaminacion', 'Medio (clase II)')}, {_n(LF_MM_KV)} mm/kV"),
           ("Velocidad de viento de diseño", "m/s", _n(CLIMA["viento_ms"], 1)), ("Radiación solar", "–", CLIMA.get("radiacion", "Muy alta (altiplano)" if ALT > 3000 else "Alta")),
           ("Nivel isoceráunico", "días/año", str(SIS["isoceraunico"]))]
ETS = [
 ("ETS-LP-01", "POSTES DE CONCRETO ARMADO CENTRIFUGADO (CAC)",
  "Fija las características mínimas, pruebas y condiciones de suministro de los postes de concreto armado centrifugado de 12 m y 13 m.",
  ["NTP 339.027:2002 Postes de hormigón (concreto) armado para líneas aéreas", "R.D. N.° 026-2003-EM/DGE", "ACI 318 – requisitos de concreto estructural"],
  ["Los postes serán troncocónicos, de sección anular, fabricados por centrifugado, con perilla de concreto en la punta, agujeros pasantes para ferretería según "
   "las láminas LD-LP y conductor de puesta a tierra embebido con terminales en la punta y a 2,0 m sobre el nivel del empotramiento.",
   "El recubrimiento mínimo del acero será de 25 mm; el concreto tendrá f'c ≥ 350 kg/cm² y relación agua/cemento ≤ 0,45; curado a vapor o húmedo.",
   "Llevarán grabado bajo relieve, a 3,0 m de la base: nombre del fabricante, fecha de fabricación, longitud/carga de trabajo/diámetros y marca de empotramiento."],
  [("Longitud total", "m", "12 / 12 / 13"), ("Carga de trabajo a 0,15 m de la punta", "daN", "300 / 400 / 400"), ("Coeficiente de seguridad", "–", "2"),
   ("Diámetro en la punta", "mm", "≥ 150 / 160 / 180"), ("Diámetro en la base", "mm", "≥ 330 / 340 / 375"), ("Carga de rotura mínima", "daN", "600 / 800 / 800"),
   ("Flecha máxima con carga de trabajo", "% L", "≤ 4 (≤ 6 % para rotura)"), ("Masa aproximada", "kg", "a declarar por el fabricante")],
  ["Ensayos de tipo: flexión hasta la rotura, carga de trabajo y deformación permanente (NTP 339.027).", "Ensayos de rutina: dimensiones, inspección visual y recubrimiento.",
   "Ensayos de aceptación: muestra según NTP 339.027 – flexión con carga de trabajo y verificación de flecha y fisuración."]),
 ("ETS-LP-02", "AISLADORES DE PORCELANA TIPO PIN CLASE ANSI 56-4",
  "Aisladores de porcelana tipo PIN para armados de suspensión y ángulo (PS1-3, PA1-3) y para soporte de puentes.",
  ["ANSI C29.1 Métodos de ensayo para aisladores", "ANSI C29.6 Aisladores de porcelana tipo PIN", "R.D. N.° 026-2003-EM/DGE"],
  ["Porcelana procesada en húmedo, esmaltada en color marrón o gris, con rosca de plomo o polietileno de 35 mm (1 3/8\") para espiga.",
   "Se suministrarán con la espiga correspondiente (cabeza de poste 609 mm o cruceta 381 mm)."],
  [("Clase ANSI", "–", "56-4"), ("Línea de fuga mínima", "mm", f"685 (requerida {fmt(K['aisl']['Lf'],0)})"), ("Tensión de flameo en seco a 60 Hz", "kV", "≥ 140"),
   ("Tensión de flameo bajo lluvia a 60 Hz", "kV", f"≥ 95 (requerida {fmt(K['aisl']['Vfi'],1)})"), ("Tensión crítica de flameo al impulso positivo", "kVp", f"≥ 225 (requerida {fmt(K['aisl']['Vi'],1)})"),
   ("Tensión crítica de flameo al impulso negativo", "kVp", "≥ 310"), ("Resistencia mecánica en voladizo", "kN", "≥ 13,3"), ("Tensión de radiointerferencia a 22 kV", "μV", "≤ 200")],
  ["Ensayos de tipo y de rutina según ANSI C29.1/C29.6.", "Ensayos de aceptación por muestreo: dimensiones, flameo en seco, porosidad, galvanizado y voladizo."]),
 ("ETS-LP-03", "AISLADORES POLIMÉRICOS DE SUSPENSIÓN 36 kV",
  "Aisladores poliméricos tipo suspensión para armados de anclaje, ángulo fuerte, retención/suspensión, derivación y terminales, y para retenidas.",
  ["IEC 61109 Aisladores compuestos para líneas aéreas", "ANSI C29.13", "IEC 62217 Ensayos de materiales poliméricos"],
  ["Núcleo de fibra de vidrio reforzada con resina, revestimiento y aletas de caucho de silicona (HTV/LSR) resistente a la radiación UV, herrajes de acero forjado "
   "galvanizado: horquilla (lado estructura) y lengüeta (lado conductor)."],
  [("Tensión nominal", "kV", "36"), ("Línea de fuga mínima", "mm", "720"), ("Tensión de sostenimiento al impulso (1,2/50 μs)", "kVp", f"≥ 170 (a {_n(ALT)} m)"),
   ("Tensión de sostenimiento a 60 Hz bajo lluvia", "kV", "≥ 100"), ("Carga mecánica especificada (SML)", "kN", "≥ 70"), ("Carga de rutina (RTL)", "kN", "≥ 35"),
   ("Longitud aproximada", "mm", "≈ 650")],
  ["Ensayos de diseño, tipo, muestreo y rutina según IEC 61109.", "Protocolos de ensayo de envejecimiento (1 000 h) y de erosión/tracking para el material de las aletas."]),
 ("ETS-LP-04", "CONDUCTOR DE ALEACIÓN DE ALUMINIO AAAC 70 mm²",
  "Conductor desnudo de aleación de aluminio para las fases de la línea primaria.",
  ["ASTM B398 Alambre de aleación de aluminio 6201-T81", "ASTM B399 Conductores cableados concéntricos AAAC", "IEC 61089", "R.D. N.° 026-2003-EM/DGE"],
  ["Conductor cableado concéntrico de 19 hilos de aleación 6201-T81, sin empalmes en los hilos durante el cableado (salvo los permitidos por ASTM B399).",
   "Se entregará en carretes de madera tratada o metálicos, con longitudes que minimicen empalmes (≥ 2 000 m por carrete)."],
  [("Sección nominal / real", "mm²", "70 / 65,8"), ("Número y diámetro de hilos", "N.° × mm", "19 × 2,10"), ("Diámetro exterior", "mm", "10,5"),
   ("Masa", "kg/km", "181"), ("Carga de rotura mínima", "kN", "19,5 (295,8 N/mm²)"), ("Módulo de elasticidad final", "N/mm²", "60 760"),
   ("Coeficiente de dilatación lineal", "1/°C", "23 × 10⁻⁶"), ("Resistencia eléctrica en c.c. a 20 °C", "Ω/km", "≤ 0,507")],
  ["Ensayos de rutina sobre hilos (diámetro, tracción, alargamiento, resistividad) y sobre el conductor (masa, cableado, resistencia).",
   "Ensayo de tensión-deformación (stress-strain) y de rotura del conductor como ensayo de tipo."]),
 ("ETS-LP-05", "ACCESORIOS DEL CONDUCTOR",
  "Varillas de armar, alambre y cinta de amarre, grapas de anclaje tipo pistola, conectores doble vía, conectores bimetálicos y manguitos de empalme.",
  ["ASTM B398/B231", "IEC 61284 Herrajes para líneas aéreas", "R.D. N.° 026-2003-EM/DGE"],
  ["Varillas de armar preformadas de aleación de aluminio para AAAC 70 mm²; alambre de amarre de Al de 16 mm²; cinta plana de Al de 1,27 × 7,62 mm.",
   "Grapas de anclaje tipo pistola de aleación de aluminio con 3 pernos U de acero galvanizado; carga de deslizamiento ≥ 90 % de la rotura del conductor.",
   "Conectores doble vía Al/Al de 3 pernos para 35–70 mm² y conectores bimetálicos Al/Cu para conexión de equipos.",
   "Manguitos de empalme a compresión con resistencia ≥ 95 % de la rotura del conductor."],
  [("Carga de deslizamiento de la grapa", "% Tr", "≥ 90"), ("Resistencia del empalme a compresión", "% Tr", "≥ 95"), ("Material de varillas y conectores", "–", "Aleación de aluminio")],
  ["Ensayos de tracción y deslizamiento, ciclos térmicos (conectores) y verificación dimensional."]),
 ("ETS-LP-06", "MATERIAL DE FERRETERÍA PARA POSTES Y CRUCETAS",
  "Crucetas y riostras de perfil angular, abrazaderas, pernos, espigas, tuercas ojo, pernos ojo, arandelas y grilletes.",
  ["ASTM A36 Acero estructural", "ASTM A123 / A153 Galvanizado en caliente", "ASTM A307 / A394 Pernos", "R.D. N.° 026-2003-EM/DGE"],
  ["Crucetas de perfil angular 64 × 64 × 6,4 × 2 400 mm (Tipo I) y riostras 64 × 64 × 6,4 × 1 000 mm, con agujeros según lámina LD-LP-12.",
   "Galvanizado en caliente después de la fabricación; masa mínima de zinc 600 g/m² (≈ 86 μm) en perfiles y 380 g/m² en pernos."],
  [("Carga de rotura de la cruceta con riostras – transversal", "kN", "≥ 12"), ("– longitudinal", "kN", "≥ 6"), ("– vertical por fase", "kN", "≥ 3"),
   ("Pernos de 16 mm: carga de rotura mínima", "kN", "≥ 55"), ("Espesor de zinc mínimo", "μm", "86")],
  ["Ensayos de galvanizado (masa y adherencia), dimensionales y de carga sobre muestras de crucetas armadas."]),
 ("ETS-LP-07", "RETENIDAS Y ANCLAJES",
  "Cable de acero, varilla de anclaje, grapas paralelas, guardacabos, abrazaderas, aislador de tracción, guardacable y bloque de anclaje.",
  ["ASTM A475 Cable de acero galvanizado grado Siemens-Martin", "ASTM A153", "R.D. N.° 026-2003-EM/DGE"],
  ["Cable de 10 mm (3/8\") de 7 hilos galvanizado clase B; varilla de anclaje de 16 mm × 2,40 m con ojal guardacabo; bloque de concreto armado de 0,50 × 0,50 × 0,20 m "
   "(f'c = 210 kg/cm², malla de 3/8\")."],
  [("Carga de rotura del cable", "kN", "≥ 30,92"), ("Carga de rotura de la varilla de anclaje", "kN", "≥ 50"), ("Carga de deslizamiento de grapas", "% Tr", "≥ 90")],
  ["Ensayos de tracción del cable y de la varilla; masa de zinc; verificación dimensional del bloque."]),
 ("ETS-LP-08", "MATERIAL PARA PUESTA A TIERRA",
  "Varillas de cobre, conductor de cobre de 25 mm², conectores, planchas dobladas, cemento conductivo, cajas de registro y tubos de protección.",
  ["UL 467 / IEEE 837 Conectores de puesta a tierra", "NTP 370.042 Conductores de cobre", "R.D. N.° 026-2003-EM/DGE"],
  ["Varilla de cobre de 16 mm × 2,40 m con rosca; conductor de cobre recocido cableado de 25 mm² (7 hilos); cemento conductivo en dosis de 25 kg."],
  [("Diámetro / longitud de la varilla", "mm / m", "16 / 2,40"), ("Sección del conductor de bajada", "mm²", "25"), ("Resistencia de puesta a tierra a obtener", "Ω", "≤ 25 (con equipos)")],
  ["Verificación de dimensiones, pureza del cobre y protocolos del cemento conductivo."]),
 ("ETS-LP-09", "SECCIONADORES FUSIBLES TIPO EXPULSIÓN (CUT-OUT) 38 kV",
  "Seccionadores fusibles unipolares para seccionamiento y protección en " + lista(["el arranque"] + [f"la derivación a {x['nombre']}" for x in DER] + [f"la llegada a {C['fin']}"]) + ".",
  ["ANSI C37.40 / C37.41 / C37.42", "IEC 60282-2"],
  ["Aislador de soporte polimérico o de porcelana, tubo portafusible con dispositivo de apertura con pértiga (load-buster), soporte para cruceta."],
  [("Tensión nominal", "kV", "38"), ("Nivel básico de aislamiento", "kVp", "170"), ("Corriente nominal", "A", "100"), ("Capacidad de interrupción (asimétrica)", "kA", "≥ 8"),
   ("Línea de fuga mínima", "mm", f"≥ 700 (requerida {fmt(K['aisl']['Lf'],0)})")],
  ["Ensayos de tipo y de rutina según ANSI C37.41."]),
 ("ETS-LP-10", "FUSIBLES TIPO EXPULSIÓN",
  f"Elementos fusibles tipo K: {fusibles_txt(C)} (selección en el Vol. IX).",
  ["ANSI C37.43", "IEC 60282-2"], ["Tipo K (rápido), con botón; curvas tiempo-corriente certificadas para la coordinación del Vol. IX."],
  [("Tipo", "–", "K"), ("Corrientes nominales", "A", "10 / 65 / 80")], ["Verificación de curvas tiempo-corriente (protocolo del fabricante)."]),
 ("ETS-LP-11", "PARARRAYOS DE ÓXIDO METÁLICO 24 kV",
  "Pararrayos para la protección de equipos en el arranque, la derivación y la llegada.",
  ["IEC 60099-4", "ANSI/IEEE C62.11"],
  ["Envolvente polimérica, desconectador automático, soporte aislante para cruceta y terminales para conductor de Cu de 25 mm²."],
  [("Tensión nominal (Ur)", "kV", "24"), ("Tensión de operación continua (MCOV)", "kV", f"≥ 19,5 (requerida {fmt(K['pr']['MCOVr'],2)})"), ("Corriente nominal de descarga", "kA", "10"),
   ("Nivel de protección 8/20 μs, 10 kA", "kVp", f"≤ 68,5"), ("Nivel de protección frente de onda", "kVp", "≤ 73,6"), ("TOV 10 s", "kV", f"≥ 25,3 (requerida {fmt(K['pr']['TOVs'],2)})"),
   ("Línea de fuga mínima", "mm", "780")],
  ["Ensayos de tipo y de rutina según IEC 60099-4 (tensión residual, TOV, descargas de larga duración)."]),
 ("ETS-LP-12", "RECONECTADOR AUTOMÁTICO TRIFÁSICO 38 kV",
  f"Reconectador para la protección y maniobra de la línea en el arranque ({C['E0']}), con su control electrónico, transformador auxiliar y soporte de poste.",
  ["IEC 62271-111 / IEEE C37.60", "IEC 60255 (relés de protección)", "IEC 60870-5-104 / DNP 3.0 (comunicaciones)"],
  ["Interrupción en vacío, aislamiento sólido o gas, mando por actuador magnético, control con funciones 50/51, 50N/51N, 51SEF, 79, registro de eventos y "
   "oscilografía; batería de respaldo ≥ 24 h; gabinete IP54; apto para 4 000 m s.n.m.", "Incluye transformador de tensión auxiliar para el control y kit de montaje en poste."],
  [("Tensión nominal", "kV", "38"), ("Nivel básico de aislamiento", "kVp", "170"), ("Corriente nominal", "A", "≥ 630"), ("Capacidad de interrupción", "kA", "≥ 12,5"),
   ("Número de operaciones mecánicas", "–", "≥ 10 000"), ("Recierres", "–", "hasta 4 operaciones a la apertura definitiva")],
  ["Ensayos de tipo según IEC 62271-111; ensayos de rutina y prueba funcional del control en fábrica con presencia de la Entidad."]),
 ("ETS-LP-13", "SEÑALIZACIÓN DE ESTRUCTURAS",
  "Placas de numeración y señales de riesgo eléctrico.",
  ["R.M. N.° 111-2013-MEM/DM", "NTP 399.010-1 Señales de seguridad"], ["Placas de plancha de acero galvanizado o aluminio con pintura esmaltada; numeración según planilla de estructuras."],
  [("Dimensiones de la placa de numeración", "mm", "120 × 80"), ("Ubicación", "m", "3,0 sobre el terreno")], ["Verificación visual y dimensional."]),
]

# ================================================================================ EETT SUMINISTRO
d = nuevo("VOLUMEN IV – ESPECIFICACIONES TÉCNICAS DE SUMINISTRO")
caratula_volumen(d, "VOLUMEN IV", "ESPECIFICACIONES TÉCNICAS DE SUMINISTRO DE MATERIALES")
p(d, "EXPEDIENTE TÉCNICO", bold=True, align="C"); p(d, f"“{NOMBRE}”", bold=True, align="C", size=11)
p(d, "CAPÍTULO III: ESPECIFICACIONES TÉCNICAS DE SUMINISTRO", bold=True, align="C")
h(d, "3.1 Generalidades")
p(d, "Las presentes especificaciones establecen los requisitos técnicos mínimos que deben cumplir los materiales y equipos de la línea primaria en 22,9 kV, "
     "su fabricación, pruebas, embalaje, transporte y entrega en el almacén de obra. Complementan a la R.D. N.° 026-2003-EM/DGE; en caso de discrepancia prevalece "
     "el requisito más exigente.")
h(d, "3.2 Alcances")
p(d, "El suministro comprende todos los materiales indicados en el metrado (Vol. V – hoja SUMINISTRO LP e INVENTARIO FÍSICO POR ESTRUCTURA), nuevos, de primera "
     "calidad y de fabricación reciente (no mayor de 12 meses a la entrega).")
h(d, "3.2.1 Ensayos y pruebas", 2)
p(d, "Los materiales serán sometidos a las pruebas de tipo, rutina y aceptación indicadas en cada especificación. Los ensayos de aceptación se efectuarán en presencia "
     "del supervisor, a cargo y costo del contratista, en laboratorios acreditados por INACAL o del fabricante con equipos calibrados.")
h(d, "3.2.2 Embalaje", 2)
p(d, f"El embalaje será adecuado para el transporte marítimo y terrestre hasta {C['fin']} (altitud {_n(C['E'].cota_m.iloc[-1], 0)} m s.n.m.), con protección contra humedad, golpes y radiación UV; "
     "cada bulto llevará marcas con el nombre del proyecto, contenido, peso bruto y neto.")
h(d, "3.2.3 Garantías", 2)
p(d, "El proveedor garantizará los materiales contra defectos de diseño, fabricación y materiales por un periodo mínimo de 24 meses desde la puesta en servicio.")
h(d, "3.2.4 Procedimientos para el control de calidad de los materiales", 2)
p(d, "El contratista presentará, antes de la fabricación, los planos, protocolos de ensayos de tipo y certificados de calidad (ISO 9001) del fabricante, y las "
     "tablas de datos técnicos garantizados llenadas y firmadas por el fabricante.")
h(d, "3.2.5 Recepción", 2)
p(d, "La recepción en el almacén de obra se hará con acta, verificando cantidad, estado, marcas y protocolos; el material observado será reemplazado sin costo.")
h(d, "3.2.6 Condiciones de operación", 2)
tabla(d, ["Condición", "Unidad", "Valor"], COND_OP, widths=[8.6, 2.4, 4.6], size=9)
for k, (cod, tit, alc, normas, cond, datos_t, pruebas) in enumerate(ETS, 3):
    salto(d)
    h(d, f"3.{k} {cod} {tit}")
    h(d, f"3.{k}.1 Alcance", 2); p(d, alc)
    h(d, f"3.{k}.2 Normas aplicables", 2); bullets(d, normas)
    h(d, f"3.{k}.3 Condiciones técnicas y de fabricación", 2)
    for x in cond: p(d, x)
    h(d, f"3.{k}.4 Pruebas", 2); bullets(d, pruebas)
    h(d, f"3.{k}.5 Información técnica requerida", 2)
    p(d, "Catálogos, planos de fabricación, protocolos de ensayos de tipo, certificado de calidad y tabla de datos técnicos garantizados llenada por el fabricante.")
    tabla(d, ["N.°", "Característica", "Unidad", "Valor requerido", "Valor garantizado"], [(i + 1, a, b, c, "") for i, (a, b, c) in enumerate(datos_t)],
          widths=[1.0, 6.4, 1.8, 3.6, 2.8], size=9, titulo=f"TABLA DE DATOS TÉCNICOS GARANTIZADOS – {tit}")
proyectista(d)
d.save(OUT / "VOL. IV ESPECIFICACIONES TECNICAS DE SUMINISTRO.docx")

# ================================================================================ EETT MONTAJE
wb = openpyxl.load_workbook(W / "presup/PRESUPUESTO.xlsx", data_only=True); mt = wb["MONTAJE LP"]
partidas = []
for r in range(1, mt.max_row + 1):
    it, ds, un = mt.cell(r, 1).value, mt.cell(r, 2).value, mt.cell(r, 3).value
    if it and ds and isinstance(ds, str) and not ds.startswith("SUB"):
        partidas.append((str(it), ds, un, mt.cell(r, 4).value))
TXT = {
 "1.01": ("Replanteo de la poligonal y ubicación de cada estructura con estación total y GPS, verificando el plantillado con las cotas reales del terreno; incluye la "
          "ingeniería de detalle (planos de replanteo, planilla definitiva y perfil corregido) para aprobación del supervisor antes de la excavación.",
          "Se mide por kilómetro de línea replanteada y aprobada."),
 "1.02": ("Estudio de coordinación de protecciones con datos reales del sistema de " + CONC + " (potencias de cortocircuito, ajustes aguas arriba) y ajuste del reconectador.",
          "Global, a la aprobación del estudio por la concesionaria."),
 "1.03": ("Monitoreo arqueológico durante las excavaciones, según el Plan de Monitoreo Arqueológico aprobado por el Ministerio de Cultura (D.S. N.° 003-2014-MC).",
          "Global, con informe final aprobado."),
 "1.04": ("Ejecución del programa de monitoreo ambiental del instrumento de gestión ambiental aprobado (aire, ruido, suelo).", "Global."),
 "1.05": ("Recolección, clasificación, almacenamiento temporal y disposición final de residuos sólidos y excedentes de obra.", "Global."),
 "1.06": ("Talleres de información a la población y autoridades sobre el proyecto, la servidumbre y las medidas de seguridad.", "Global."),
 "1.07": ("Suministro e instalación del cartel de obra de 3,60 × 2,40 m según el modelo de la Entidad.", "Unidad instalada."),
}
GEN = {"2": ("Carga, transporte y descarga desde el almacén de obra hasta el punto de instalación, con equipo adecuado que evite daños a los materiales.", "Por unidad, juego o kilómetro transportado."),
       "3": ("Excavación con las dimensiones del Anexo 5 y lámina LD-LP-11; izaje con camión grúa, alineamiento y verticalidad (desviación ≤ 1/250); solado y relleno con "
             "concreto ciclópeo; eliminación de excedentes. Incluye pintado de numeración y señalización.", "Por metro cúbico (excavación, eliminación) o por unidad (izaje, cimentación, solado)."),
       "4": ("Excavación de 0,80 × 0,80 × 2,00 m en la dirección del cable, instalación de la varilla y el bloque de anclaje, relleno compactado en capas de 0,20 m, montaje del "
             "cable, aislador de tracción, guardacable y tensado (lámina LD-LP-08).", "Por unidad instalada o metro cúbico."),
       "5": ("Montaje de crucetas, riostras, abrazaderas, aisladores y accesorios según las láminas LD-LP; ajuste de pernos con torquímetro; montaje de equipos con sus "
             "conexiones y puesta a tierra.", "Por juego (armado) o unidad (equipo) instalado y aprobado."),
       "6": ("Tendido con poleas y freno, sin arrastre del conductor sobre el terreno; flechado según la tabla de templado (Anexo 4) a la temperatura del conductor medida "
             "con termómetro de contacto; amarre en aisladores PIN con varillas de armar y en anclajes con grapas pistola. Empalmes a compresión fuera de los vanos de cruce.",
             "Por kilómetro de conductor por fase tendido y flechado."),
       "7": ("Excavación del pozo, instalación de la varilla, relleno con tierra de cultivo cernida y cemento conductivo, conexión de la bajada de cobre y medición de la "
             "resistencia (≤ 25 Ω en estructuras con equipos).", "Por unidad instalada con protocolo de medición o metro cúbico."),
       "8": ("Pruebas de continuidad, secuencia de fases, aislamiento (megóhmetro 5 kV), resistencia de puesta a tierra, inspección final y puesta en servicio coordinada "
             "con " + CONC + "; elaboración del expediente conforme a obra.", "Por kilómetro probado o global.")}
d = nuevo("VOLUMEN IV – ESPECIFICACIONES TÉCNICAS DE MONTAJE")
caratula_volumen(d, "VOLUMEN IV", "ESPECIFICACIONES TÉCNICAS DE MONTAJE ELECTROMECÁNICO")
p(d, "EXPEDIENTE TÉCNICO", bold=True, align="C"); p(d, f"“{NOMBRE}”", bold=True, align="C", size=11)
p(d, "CAPÍTULO IV: ESPECIFICACIONES TÉCNICAS DE MONTAJE", bold=True, align="C")
sec = [("4.1 De la programación", [("4.1.1 Cronograma de ejecución", "El contratista presentará, dentro de los 7 días de firmado el contrato, el cronograma de ejecución (Gantt / PERT-CPM) "
                                     "concordante con el plazo de " + str(PROYECTO["plazo_dias"]) + " días y con el calendario valorizado del Vol. V."),
                                    ("4.1.2 Plazos contractuales", "El plazo se computa desde el día siguiente de cumplidas las condiciones de inicio de obra; las ampliaciones se tramitarán según la Ley de Contrataciones del Estado y su reglamento."),
                                    ("4.1.3 Modificación del cronograma", "Toda modificación requerirá la aprobación del supervisor y no implicará ampliación de plazo salvo causales reconocidas."),
                                    ("4.1.4 Cuaderno de obra", "Se llevará un cuaderno de obra legalizado donde el residente y el supervisor anotarán las ocurrencias, consultas y órdenes.")]),
       ("4.2 Del personal", [("4.2.1 Personal clave", "Ingeniero residente (electricista o mecánico electricista colegiado y habilitado) y especialista en seguridad y salud en el trabajo, según el Vol. X."),
                             ("4.2.2 Desempeño y leyes sociales", "El contratista cumplirá las leyes laborales y de construcción civil y retirará al personal que el supervisor objete fundadamente."),
                             ("4.2.3 Seguridad e higiene", "Se aplicará el Reglamento de Seguridad y Salud en el Trabajo con Electricidad (R.M. N.° 111-2013-MEM/DM), con plan de seguridad, "
                              "análisis de trabajo seguro, equipos de protección personal, revelador de tensión y puestas a tierra temporales.")]),
       ("4.3 De la ejecución", [("4.3.1 Ejecución de los trabajos", "Los trabajos se ejecutarán según los planos y láminas del Vol. VII, estas especificaciones y la R.D. N.° 029-2003-EM/DGE."),
                                ("4.3.2 Herramientas y equipos", "El contratista dispondrá de camión grúa, winche, freno, poleas, dinamómetro, estación total, telurómetro y megóhmetro calibrados."),
                                ("4.3.3 Cambios y rechazos", "Los cambios requerirán aprobación escrita; los trabajos defectuosos serán rehechos sin costo para la Entidad."),
                                ("4.3.4 Daños y protección del medio ambiente", "El contratista repondrá los daños a terceros y cumplirá las medidas del instrumento de gestión ambiental; restituirá las áreas intervenidas."),
                                ("4.3.5 Vigilancia y limpieza", "Mantendrá vigilancia de almacenes y frentes de trabajo y dejará la zona limpia al término de cada frente.")]),
       ("4.4 De la supervisión", [("4.4.1 Supervisión de la obra", "La Entidad designará un supervisor con las calificaciones del Vol. X, quien controlará la calidad, el avance y las valorizaciones."),
                                  ("4.4.2 Facilidades de inspección", "El contratista facilitará el acceso a almacenes, fábricas y frentes de trabajo.")]),
       ("4.5 De la aceptación", [("4.5.1 Procedimiento general", "Concluida la obra, el comité de recepción verificará las instalaciones con el expediente conforme a obra."),
                                 ("4.5.2 Pruebas en blanco", "Inspección visual, verificación de distancias de seguridad y flechas, continuidad, aislamiento y resistencia de puesta a tierra."),
                                 ("4.5.3 Prueba de puesta en servicio", "Energización coordinada con " + CONC + ", medición de tensiones y secuencia de fases, operación del reconectador; 72 h de operación sin falla."),
                                 ("4.5.4 Garantía", "El contratista garantizará la obra por 12 meses desde la recepción, corrigiendo los defectos que se presenten.")])]
for t, items in sec:
    h(d, t)
    for st, tx in items:
        h(d, st, 2); p(d, tx)
h(d, "4.6 Especificaciones particulares por partida")
for it, ds, un, met in partidas:
    if "." not in it:
        h(d, f"{it}.0 {ds}", 2); continue
    h(d, f"{it} {ds}", 3)
    desc, med = TXT.get(it, GEN.get(it.split(".")[0], ("", "")))
    pr(d, [("Descripción y ejecución: ", True), (desc, False)])
    pr(d, [("Unidad de medida: ", True), (f"{un}. ", False), ("Método de medición: ", True), (med, False)])
    pr(d, [("Bases de pago: ", True), ("la cantidad medida se pagará al precio unitario del contrato, que comprende mano de obra, equipos, herramientas, materiales de "
                                       "montaje, leyes sociales e imprevistos necesarios para completar la partida.", False)])
proyectista(d)
d.save(OUT / "VOL. IV ESPECIFICACIONES TECNICAS DE MONTAJE.docx")
print("ok s13b")
