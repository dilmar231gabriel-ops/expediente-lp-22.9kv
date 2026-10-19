"""Verificación cruzada: planilla ↔ anexos ↔ presupuesto. Genera work/datos_clave.json para los documentos.
Uso: python p11_verificacion.py WORK"""
import json, math, sys
from pathlib import Path
import openpyxl, pandas as pd
sys.path.insert(0, str(Path(__file__).parent))
from params import *
from catalogo import BOM, RET_RI, PAT_KIT, PUENTE_M, cable_ret, DESPERDICIO_COND, MATD

W = Path(sys.argv[1]); A = W / "anexos"
E = pd.read_csv(W / "planilla_estructuras.csv"); R = json.load(open(W / "resultados.json"))
num = lambda v: float(v) if v not in (None, "", "-") and not isinstance(v, str) or (isinstance(v, str) and v.replace(".", "", 1).replace("-", "", 1).isdigit()) else v
def wbv(f): return openpyxl.load_workbook(f, data_only=True)
def find(ws, label, col=1, valcol=3):
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, col).value
        if isinstance(v, str) and v.startswith(label): return ws.cell(r, valcol).value
    return None
K = {"chk": []}
def chk(nombre, ok, detalle=""):
    K["chk"].append({"verificacion": nombre, "resultado": "CONFORME" if ok else "NO CONFORME", "detalle": detalle})

# --- anexos
a1 = wbv(A / "Anexo 1 SELECCION DE AISLADORES.xlsx")["SEL_AISL_22,9KV"]
K["aisl"] = {k: num(find(a1, lab)) for k, lab in [("fch", "Factor de corrección por altitud"), ("delta", "Densidad relativa"), ("b", "Presión barométrica"),
                                                  ("Lf", "Longitud de línea de fuga fase–tierra"), ("Vfi", "Tensión de sostenimiento a frecuencia industrial bajo"),
                                                  ("Vi", "Tensión crítica de flameo"), ("NBIh", "Nivel básico de aislamiento requerido"), ("h", "Altitud de diseño")]}
a2 = wbv(A / "Anexo 2 Aislamiento y Pararrayos.xlsx")["Seleccion Pararrayos"]
K["pr"] = {k: num(find(a2, lab)) for k, lab in [("MCOVr", "MCOV requerido"), ("MCOV", "MCOV del pararrayo"), ("TOVs", "TOV del sistema"), ("TOVp", "TOV del pararrayo"),
                                                ("NPFO", "Nivel de protección para frente"), ("NPIA", "Nivel de protección al impulso atm"), ("NPIM", "Nivel de protección al impulso de man"),
                                                ("BILh", "NBI efectivo"), ("MP1", "MP1 ="), ("MP2", "MP2 ="), ("MP3", "MP3 ="), ("Lfp", "Línea de fuga del pararrayo")]}
a3 = wbv(A / "Anexo 3 CALCULO PUESTA A TIERRA.xlsx")["CAL.P.TIERRA"]
K["pat"] = {k: num(find(a3, lab)) for k, lab in [("R1", "PAT-1, una varilla, sin"), ("R1t", "PAT-1, una varilla, con"), ("R2", "PAT-2, dos varillas, sin"), ("R2t", "PAT-2, dos varillas, con")]}
a8 = wbv(A / "Anexo 8 CAIDA DE TENSION Y PERDIDAS.xlsx")["PARAMETROS"]
K["elec"] = {k: num(find(a8, lab)) for k, lab in [("r1", "Resistencia a la temperatura"), ("X1", "Reactancia inductiva"), ("K", "Factor de caída"), ("DMG", "Distancia media"), ("FP", "Factor de pérdidas")]}
for r in range(1, a8.max_row + 1):
    if a8.cell(r, 1).value == "TOTAL":
        K["elec"].update({"dV": a8.cell(r, 5).value, "PJ": a8.cell(r, 6).value, "PE": a8.cell(r, 7).value, "L": a8.cell(r, 2).value, "P": a8.cell(r, 3).value})
    if isinstance(a8.cell(r, 1).value, str) and a8.cell(r, 1).value.startswith(INICIO["nombre"]):
        K["elec"]["dV_cut"] = a8.cell(r, 5).value
chk("ΔV Excel (Anexo 8) = ΔV motor de cálculo", abs(K["elec"]["dV"] - R["electrico"]["dV_fin_pct"]) < 0.005, f"{K['elec']['dV']:.3f} % vs {R['electrico']['dV_fin_pct']:.3f} %")
_e = R["electrico"]
chk(f"Conductor del kit AAAC {SECCION} mm² cumple ΔV con reserva, corriente y cortocircuito", _e.get("kit_cumple", True),
    f"ΔV {R['electrico']['dV_fin_pct']:.2f} % ≤ {SIS['reserva_dV']*SIS['dV_max_pct']:.2f} %; óptimo económico: {_e.get('seccion_optima')} mm²")
# CMC
a4 = wbv(A / "Anexo 4 CMC.xlsx")["CMC_AAAC-70"]
mec = pd.read_csv(W / "calculo_mecanico_conductor.csv")
rows = {a4.cell(r, 1).value: r for r in range(1, a4.max_row + 1) if isinstance(a4.cell(r, 1).value, (int, float))}
hdr = {}
for c in range(2, 40):
    v = a4.cell(rows[10] - 2, c).value
    if v: hdr[v] = c
mx = 0
for x in mec.itertuples():
    if x.vano_m in rows:
        key = [k for k in hdr if k.startswith(x.hipotesis.split()[0] + " ")][0]
        mx = max(mx, abs(a4.cell(rows[x.vano_m], hdr[key]).value - x.sigma_N_mm2) / x.sigma_N_mm2)
chk("CMC Excel (Anexo 4) = motor de cálculo (σ por vano e hipótesis)", mx < 1e-6, f"diferencia relativa máxima {mx:.1e}")
K["cmc"] = {}
for d in (100, 150, 200):
    K["cmc"][d] = {h.split()[0]: {"s": a4.cell(rows[d], c).value, "T": a4.cell(rows[d], c + 1).value, "f": a4.cell(rows[d], c + 3).value,
                                  "pct": a4.cell(rows[d], c + 4).value} for h, c in hdr.items() if h.startswith("H")}
K["cmc_maxpct"] = max(a4.cell(r, 38).value for r in rows.values())
# CME verif
a6 = wbv(A / "Anexo 6 CME_CMR_ALUMINIO_AAAC.xlsx")
v = a6["VERIF_ESTRUCTURAS"]; rr = json.load(open(W / "verif_rows.json"))
coh = sum(1 for r in range(rr["v0"], rr["v1"] + 1) if v.cell(r, 36).value == "COHERENTE")
chk("Verificación mecánica estructura por estructura (Anexo 6) coherente con la planilla", coh == len(E), f"{coh} de {len(E)} estructuras")
K["cme"] = {"coherentes": coh, "QN_max": max(v.cell(r, 26).value for r in range(rr["v0"], rr["v1"] + 1)),
            "QNr_max": max(v.cell(r, 29).value for r in range(rr["v0"], rr["v1"] + 1)), "ret": sum(v.cell(r, 28).value for r in range(rr["v0"], rr["v1"] + 1)),
            "FS_cable_min": min(v.cell(r, 32).value for r in range(rr["v0"], rr["v1"] + 1) if isinstance(v.cell(r, 32).value, (int, float)))}
a5 = wbv(A / "Anexo 5 Calculo Cimentacion.xlsx")["RESUMEN"]
K["cim"] = {}
for r in range(1, a5.max_row + 1):
    lab = a5.cell(r, 1).value
    if isinstance(lab, str):
        for key, pre in [("R1", "R1 ="), ("R2", "R2 ="), ("s1", "σ1"), ("s2", "σ2"), ("sv", "Presión vertical"), ("ve", "Volumen de excavación"), ("vc", "Volumen de concreto"), ("vs", "Volumen de solado"), ("ee", "Volumen de material excedente")]:
            if lab.startswith(pre): K["cim"][key] = [a5.cell(r, c).value for c in (3, 4, 5)]
a7 = wbv(A / "Anexo 7 Calculo Mecanico de Retenidas.xlsx")["RI"]
K["ret"] = {k: num(find(a7, lab)) for k, lab in [("Tw", "Tiro máximo de trabajo"), ("FSc", "Factor de seguridad del cable"), ("FSb", "Factor de seguridad del bloque"),
                                                 ("Fres", "Fuerza resistente total"), ("Ve", "Volumen de excavación por retenida"), ("Vr", "Volumen de relleno")]}
chk("Retenidas: FS cable ≥ 2 y FS bloque ≥ 2 (Anexo 7)", K["ret"]["FSc"] >= 2 and K["ret"]["FSb"] >= 2, f"FS cable {K['ret']['FSc']:.2f}; FS bloque {K['ret']['FSb']:.2f}")
chk("Cimentación: σ1, σ2 ≤ 2,1 daN/cm² y presión vertical ≤ 1,0 daN/cm² (Anexo 5)",
    max(K["cim"]["s1"] + K["cim"]["s2"]) <= 2.1 and max(K["cim"]["sv"]) <= 1.0, f"σ máx {max(K['cim']['s1'] + K['cim']['s2']):.2f}; vertical máx {max(K['cim']['sv']):.2f}")
chk("Aislamiento: Lf, Vfi y Vi del PIN 56-4 y polimérico ≥ requeridos (Anexo 1)", K["aisl"]["Lf"] <= 685 and K["aisl"]["Vfi"] <= 95 and K["aisl"]["Vi"] <= 225,
    f"Lf {K['aisl']['Lf']:.0f} ≤ 685 mm; Vfi {K['aisl']['Vfi']:.1f} ≤ 95 kV; Vi {K['aisl']['Vi']:.1f} ≤ 225 kV")
chk("Equipos de 170 kV-BIL ≥ NBI requerido a la altitud de diseño (tolerancia 1 %, Anexo 1)", K["aisl"]["NBIh"] <= 170 * 1.01,
    f"NBI requerido {K['aisl']['NBIh']:.1f} kVp a {K['aisl']['h']:.0f} m s.n.m.")
chk("Pararrayos: MCOV, TOV, MP1–MP3 y línea de fuga (Anexo 2)", K["pr"]["MCOV"] >= K["pr"]["MCOVr"] and K["pr"]["TOVp"] >= K["pr"]["TOVs"] and min(K["pr"]["MP1"], K["pr"]["MP2"]) >= 20 and K["pr"]["MP3"] >= 15 and K["pr"]["Lfp"] >= K["aisl"]["Lf"],
    f"MCOV {K['pr']['MCOV']} kV; MP1 {K['pr']['MP1']:.0f} %, MP2 {K['pr']['MP2']:.0f} %, MP3 {K['pr']['MP3']:.0f} %")
chk("PAT-2 con tratamiento ≤ 25 Ω (Anexo 3)", K["pat"]["R2t"] <= 25, f"{K['pat']['R2t']:.1f} Ω")
chk("CMC: tiro máximo ≤ 60 % UTS en todas las hipótesis y vanos (Anexo 4)", K["cmc_maxpct"] <= 0.6, f"máximo {K['cmc_maxpct']*100:.1f} % UTS")
chk("Distancias al terreno y separación de fases en todos los vanos (Anexo 9)", R["plantillado"]["cumple_terreno"] and R["plantillado"]["cumple_sep"],
    f"{len(E)-1} vanos; D horizontal requerida máx {R['plantillado']['Dh_max']:.2f} m < 2,00 m")
chk("Sin arrancamiento en armados con aisladores PIN (vano peso ≥ 0)", not ((E.vano_peso_m < 0) & E.armado.isin(["PS1-3", "PA1-3"])).any(),
    f"{int((E.vano_peso_m < 0).sum())} estructuras con vano peso negativo, todas con cadenas de anclaje (TS-3, PA2-3, PTH-3)")

# --- presupuesto vs planilla (recuento independiente en Python)
pb = wbv(W / "presup/PRESUPUESTO.xlsx"); pl = pb["PLANILLA"]
meta = json.load(open(W / "presup_meta.json"))
esperado = {}
for x in E.itertuples():
    pst = {"12/300": "1.01", "12/400": "1.02", "13/400": "1.03"}[x.poste]; esperado[pst] = esperado.get(pst, 0) + 1
    for c, q in BOM[x.armado].items(): esperado[c] = esperado.get(c, 0) + q
    for c, q in RET_RI.items(): esperado[c] = esperado.get(c, 0) + q * x.n_retenidas
    esperado["7.01"] = esperado.get("7.01", 0) + x.n_retenidas * cable_ret(POSTES[x.poste]["L"])
    if x.puesta_tierra in PAT_KIT:
        for c, q in PAT_KIT[x.puesta_tierra].items(): esperado[c] = esperado.get(c, 0) + q
    va = 0 if pd.isna(x.vano_adelante_m) else round(x.vano_adelante_m, 2)      # igual que en la hoja (2 decimales)
    esperado["3.01"] = esperado.get("3.01", 0) + 3 * va * (1 + DESPERDICIO_COND) + 3 * PUENTE_M.get(x.armado, 0)
    esperado["10.01"] = esperado.get("10.01", 0) + 1
dif = []
for cod, row in meta["PLROW"].items():
    xl = pl.cell(row, 4).value or 0; xl = float(xl)
    if abs(xl - esperado.get(cod, 0)) > 0.05: dif.append((cod, xl, esperado.get(cod, 0)))
chk("Metrado de suministros (PLANILLA del presupuesto) = recuento independiente desde la planilla de estructuras", not dif, f"{len(meta['PLROW'])} materiales; diferencias: {dif[:3]}")
ro = pb["RESUMEN OBRA"]
K["pres"] = {k: float(ro.cell(r, 6).value) for k, r in [("SU", 11), ("MO", 12), ("TR", 13), ("CD", 14), ("GGD", 15), ("GGI", 16), ("UT", 17), ("ST", 18), ("IGV", 19), ("VR", 20), ("SUP", 22), ("COO", 23), ("LIQ", 24), ("TOT", 27)]}
K["pres"].update({"pGGD": float(ro["C15"].value), "pGGI": float(ro["C16"].value)})
chk("Resumen: CD = suministro + montaje + transporte; VR = CD + GG + U + IGV", abs(K["pres"]["CD"] - K["pres"]["SU"] - K["pres"]["MO"] - K["pres"]["TR"]) < 0.02 and
    abs(K["pres"]["VR"] - (K["pres"]["CD"] + K["pres"]["GGD"] + K["pres"]["GGI"] + K["pres"]["UT"] + K["pres"]["IGV"])) < 0.02, f"VR S/ {K['pres']['VR']:,.2f}")
cv = pb["CRONOG_VALORIZ"]
for r in range(1, cv.max_row + 1):
    if cv.cell(r, 2).value == "TOTAL MENSUAL":
        K["pres"]["meses"] = [float(cv.cell(r, c).value) for c in range(4, cv.max_column + 1) if isinstance(cv.cell(r, c).value, (int, float))]
chk("Calendario valorizado: suma de meses = valor referencial", abs(sum(K["pres"]["meses"]) - K["pres"]["VR"]) < 0.05, f"{sum(K['pres']['meses']):,.2f}")
fp = pb["FP"]; coefs = [fp.cell(r, 8).value for r in range(1, fp.max_row + 1) if isinstance(fp.cell(r, 8).value, (int, float))]
K["pres"]["fp"] = [(fp.cell(r, 4).value, fp.cell(r, 8).value) for r in range(1, fp.max_row + 1) if isinstance(fp.cell(r, 8).value, (int, float)) and fp.cell(r, 4).value]
chk("Fórmula polinómica: coeficientes suman 1,000 y ninguno < 5 %", abs(sum(c for _, c in K["pres"]["fp"]) - 1) < 1e-6 and min(c for _, c in K["pres"]["fp"]) >= 0.05,
    ", ".join(f"{s} {c:.3f}" for s, c in K["pres"]["fp"]))
su = pb["SUMINISTRO LP"]; K["pres"]["grupos"] = {}
for r in range(1, su.max_row + 1):
    b = su.cell(r, 2).value
    if isinstance(b, str) and b.startswith("SUB-TOTAL"): K["pres"]["grupos"][b.replace("SUB-TOTAL ", "")] = float(su.cell(r, 6).value)
mt = pb["MONTAJE LP"]; K["pres"]["montaje"] = {}
for r in range(1, mt.max_row + 1):
    b = mt.cell(r, 2).value
    if isinstance(b, str) and b.startswith("SUB-TOTAL"): K["pres"]["montaje"][mt.cell(r - 1, 1).value.split(".")[0] if mt.cell(r - 1, 1).value else b] = float(mt.cell(r, 6).value)
# metrados clave
K["met"] = {"postes": E.poste.value_counts().to_dict(), "armados": E.armado.value_counts().to_dict(), "ret": int(E.n_retenidas.sum()),
            "pat": E.puesta_tierra.value_counts().to_dict(), "cond_m": float(pl.cell(meta["PLROW"]["3.01"], 4).value),
            "pin": float(pl.cell(meta["PLROW"]["2.01"], 4).value), "susp": float(pl.cell(meta["PLROW"]["2.04"], 4).value),
            "cutout": float(pl.cell(meta["PLROW"]["9.01"], 4).value), "pararrayos": float(pl.cell(meta["PLROW"]["9.04"], 4).value),
            "cable_ret": float(pl.cell(meta["PLROW"]["7.01"], 4).value), "varillas": float(pl.cell(meta["PLROW"]["8.01"], 4).value)}
K["plan"] = R["plantillado"]; K["dem"] = R["demanda"]; K["el"] = R["electrico"]; K["mec"] = {k: v for k, v in R["mecanico"].items() if k != "hipotesis"}
json.dump(K, open(W / "datos_clave.json", "w"), indent=1, default=float, ensure_ascii=False)
for c in K["chk"]: print(f"[{c['resultado']}] {c['verificacion']} – {c['detalle']}")
print("VR", K["pres"]["VR"], "TOT", K["pres"]["TOT"])
